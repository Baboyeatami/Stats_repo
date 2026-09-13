"""Independently audit saved predictions, splits, metrics and documents.

Run after the build: python verify_materials.py
Optional PDF contact sheet: python verify_materials.py --render (needs pypdfium2).
"""
from pathlib import Path
import argparse
import hashlib
import json

import nbformat
import numpy as np
import pandas as pd
from docx import Document
from pypdf import PdfReader
from sklearn.metrics import accuracy_score, confusion_matrix, precision_recall_fscore_support

from build_materials import STEPS, AUTHOR

ROOT = Path(__file__).resolve().parent


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    out = ROOT / "results"
    raw = pd.read_csv(ROOT / "Occupancy_Estimation.csv")
    raw["source_row"] = raw.index + 2
    raw["timestamp"] = pd.to_datetime(raw.Date + " " + raw.Time)
    raw = raw.set_index("source_row")
    assignments = pd.read_csv(out / "split_assignments.csv")
    predictions = pd.read_csv(out / "predictions.csv")
    metrics = pd.read_csv(out / "metrics.csv")
    matrices = json.loads((out / "confusion_matrices.json").read_text())
    reports = json.loads((out / "classification_reports.json").read_text())
    assert len(metrics) == 48
    configs = {"All features", "Without Temperature", "Without Light", "Without Sound", "Without CO2", "Without PIR"}
    dates = sorted(raw.timestamp.dt.date.unique())
    splits = {"random_80_20", "temporal_holdout"} | {f"daily_{day}" for day in dates[1:]}
    expected_pairs = {(s, c) for s in splits for c in configs}
    assert len(expected_pairs) == 48
    assert not metrics.duplicated(["split", "configuration"]).any()
    assert set(zip(metrics.split, metrics.configuration)) == expected_pairs
    assert set(zip(predictions.split, predictions.configuration)) == expected_pairs
    assert set(assignments.split) == splits
    assert set(assignments.role) == {"train", "test"}
    for artifact in [reports, matrices]:
        assert set(artifact) == splits
        assert {(s, c) for s, values in artifact.items() for c in values} == expected_pairs
    assert not assignments.duplicated(["split", "source_row"]).any()
    checked = 0
    for split, group in assignments.groupby("split"):
        train = group.loc[group.role == "train", "source_row"]
        test = group.loc[group.role == "test", "source_row"]
        assert not set(train) & set(test)
        if split != "random_80_20":
            assert raw.loc[train, "timestamp"].max() < raw.loc[test, "timestamp"].min()
        if split.startswith("daily_"):
            assert raw.loc[test, "timestamp"].min() - raw.loc[train, "timestamp"].max() > pd.Timedelta(minutes=10)
        for config, pred in predictions[predictions.split == split].groupby("configuration"):
            assert not pred.source_row.duplicated().any()
            assert set(pred.source_row) == set(test)
            assert np.array_equal(raw.loc[pred.source_row, "Room_Occupancy_Count"], pred.actual)
            row = metrics[(metrics.split == split) & (metrics.configuration == config)].iloc[0]
            assert np.isclose(row.accuracy, accuracy_score(pred.actual, pred.predicted))
            assert row.errors == (pred.actual != pred.predicted).sum()
            cm = confusion_matrix(pred.actual, pred.predicted, labels=[0, 1, 2, 3])
            assert np.array_equal(cm, matrices[split][config])
            for average, labels, suffix in [("weighted", None, "weighted"),
                                             ("macro", sorted(pred.actual.unique()), "macro_observed")]:
                values = precision_recall_fscore_support(pred.actual, pred.predicted, average=average, labels=labels, zero_division=0)
                for metric, value in zip(["precision", "recall", "f1"], values):
                    assert np.isclose(row[f"{metric}_{suffix}"], value)
            p, r, f, support = precision_recall_fscore_support(pred.actual, pred.predicted, labels=[0, 1, 2, 3], zero_division=0)
            assert np.isclose(row.f1_macro_fixed4, np.mean(f))
            for c in range(4):
                report = reports[split][config][str(c)]
                assert np.allclose([report[k] for k in ["precision", "recall", "f1-score", "support"]], [p[c], r[c], f[c], support[c]])
            assert np.isclose(reports[split][config]["macro avg fixed4"]["f1-score"], row.f1_macro_fixed4)
            assert np.isclose(reports[split][config]["macro avg observed"]["f1-score"], row.f1_macro_observed)
            baseline = metrics[(metrics.split == split) & (metrics.configuration == "All features")].iloc[0]
            assert np.isclose(row.accuracy_drop_pp, 100 * (baseline.accuracy - row.accuracy))
            assert np.isclose(row.macro_f1_drop_pp, 100 * (baseline.f1_macro_observed - row.f1_macro_observed))
            checked += 1
    assert checked == len(expected_pairs)
    importance = pd.read_csv(out / "feature_importance.csv")
    expected = importance.groupby(["split", "group"], as_index=False).importance.sum()
    pd.testing.assert_frame_equal(expected, pd.read_csv(out / "group_importance.csv"), atol=1e-14)
    assert np.allclose(expected.groupby("split").importance.sum(), 1)
    repeat = json.loads((out / "repeatability_check.json").read_text())
    assert repeat["status"] == "PASS"
    for name, digest in repeat["matching_sha256"].items():
        assert hashlib.sha256((out / name).read_bytes()).hexdigest() == digest
    nb = nbformat.read(ROOT / "Room_Occupancy_Estimation.ipynb", as_version=4)
    nbformat.validate(nb)
    cells = [c for c in nb.cells if c.cell_type == "code"]
    assert len(cells) == len(STEPS) == 12
    for cell, (_, _, code) in zip(cells, STEPS):
        assert cell.source == code and cell.execution_count is not None
        assert not any(o.output_type == "error" for o in cell.outputs)
    pdf = PdfReader(ROOT / "Room_Occupancy_Student_Guide.pdf")
    pdf_text = "\n".join(p.extract_text() for p in pdf.pages)
    doc = Document(ROOT / "Room_Occupancy_Student_Guide.docx")
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert AUTHOR in pdf_text and AUTHOR in doc_text
    for title, _, _ in STEPS:
        assert title in doc_text and title in pdf_text
    headers = ["Configuration", "Accuracy", "Weighted F1", "Macro F1 observed", "Macro drop pp"]
    score_tables = [t for t in doc.tables if [c.text for c in t.rows[0].cells] == headers]
    assert len(score_tables) == 2
    section_start = pdf_text.rindex(STEPS[6][0])
    section_end = pdf_text.index(STEPS[7][0], section_start)
    score_section = pdf_text[section_start:section_end]
    for number, split in enumerate(["random_80_20", "temporal_holdout"]):
        section = score_section[score_section.index(split):]
        if split == "random_80_20":
            section = section[:section.index("temporal_holdout")]
        normalized = " ".join(section.split())
        expected_rows = []
        for _, row in metrics[metrics.split == split].iterrows():
            values = [row.configuration, f"{row.accuracy:.4f}", f"{row.f1_weighted:.4f}",
                      f"{row.f1_macro_observed:.4f}", f"{row.macro_f1_drop_pp:+.3f}"]
            expected_rows.append(values)
            assert " ".join(values) in normalized, (split, values)
        assert [[c.text for c in r.cells] for r in score_tables[number].rows[1:]] == expected_rows
    if args.render:
        import pypdfium2 as pdfium
        from PIL import Image, ImageDraw
        document = pdfium.PdfDocument(ROOT / "Room_Occupancy_Student_Guide.pdf")
        thumbs = []
        for i in range(len(document)):
            page = document[i]
            bitmap = page.render(scale=1)
            picture = bitmap.to_pil().copy()
            picture.thumbnail((306, 396))
            thumbs.append(picture)
            bitmap.close()
            page.close()
        sheet = Image.new("RGB", (4 * 326, ((len(thumbs) + 3) // 4) * 426), "#cccccc")
        draw = ImageDraw.Draw(sheet)
        for i, picture in enumerate(thumbs):
            x, y = (i % 4) * 326 + 10, (i // 4) * 426 + 20
            sheet.paste(picture, (x, y))
            draw.text((x, y - 15), f"Page {i + 1}", fill="black")
        sheet.save(out / "pdf_contact_sheet.png")
        document.close()
    result = {"status": "PASS", "experiments_checked": checked, "code_cells_checked": len(cells),
              "pdf_pages": len(pdf.pages), "checks": ["predictions versus original labels", "all metrics and drops",
              "all matrices and per-class reports", "temporal splits and embargoes", "group importance sums",
              "repeatability artifact hashes", "notebook source/execution", "guide sections and attribution",
              "DOCX/PDF score table rows matched by split and configuration"]}
    (out / "verification_report.json").write_text(json.dumps(result, indent=2))
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
