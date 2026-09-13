"""Generate an executed notebook and matching student-guide DOCX/PDF.

python build_materials.py --verify
python build_materials.py --documents-only
"""
from pathlib import Path
import argparse
import hashlib
import json
import sys
import textwrap
from xml.sax.saxutils import escape

import nbformat
from nbclient import NotebookClient
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether

ROOT = Path(__file__).resolve().parent
AUTHOR = "Prepared by Engr. Jamie Eduardo Rosal, MSCpE"
TITLE = "Room Occupancy Estimation: Student Guide"
STEPS = []


def step(title, explanation, code):
    STEPS.append((title, textwrap.dedent(explanation).strip(), textwrap.dedent(code).strip()))


step("1. Understand the objectives", """
Goal: classify whether a room contains 0, 1, 2, or 3 people using environmental measurements. This is multiclass classification, not continuous regression.
Objective 1: develop the full-feature Random Forest. Objective 2: remove each of five sensor groups and retrain. Objective 3: compare ablation effects with individual and grouped feature importance.
A Random Forest combines decision trees. Ablation means removing one group, retraining, and comparing on exactly the same test rows. The proposal title is Random Forest-Based Sensor Ablation Analysis for Environmental Feature Contribution in Multiclass Room Occupancy Estimation.
Checkpoint: explain why the occupancy label must never be an input. This dataset cannot establish performance on occupancy counts outside 0-3.
""", """
print("Objectives: classify occupancy; compare five ablations; assess sensor contribution.")
""")

step("2. Prepare the environment", """
Materials: laptop/desktop, Python 3.12, JupyterLab, and the supplied Occupancy_Estimation.csv. A practical recommendation is 8 GB RAM, not a measured minimum. Internet is needed for installation only.
Open a terminal in this project folder. macOS/Linux: run python3.12 -m venv .venv and then source .venv/bin/activate. Windows Command Prompt: run py -3.12 -m venv .venv and then .venv\\Scripts\\activate.bat.
Run python -m pip install -r requirements.txt. Register the environment with python -m ipykernel install --user --name occupancy-guide --display-name "Occupancy Guide". Launch python -m jupyterlab, open Room_Occupancy_Estimation.ipynb, and select the Occupancy Guide kernel.
Use the same Python for installation and execution. This avoids the system-Python versus Anaconda mismatch encountered in the initial build. The local CSV makes ucimlrepo unnecessary.
Action: run the cell. Expected: Python executable and package versions. Checkpoint: resolve import errors before continuing; do not suppress warnings globally.
""", """
from pathlib import Path
import sys, json, hashlib, importlib.metadata as metadata
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from IPython.display import display
from sklearn.ensemble import RandomForestClassifier
from sklearn.pipeline import make_pipeline
from sklearn.impute import SimpleImputer
from sklearn.model_selection import train_test_split
from sklearn.metrics import (accuracy_score, precision_recall_fscore_support,
                             classification_report, confusion_matrix, ConfusionMatrixDisplay)
SEED, LABELS = 42, [0, 1, 2, 3]
DATA = Path("Occupancy_Estimation.csv")
assert DATA.exists(), "Start Jupyter in the project folder containing the CSV."
OUT = Path("results")
OUT.mkdir(exist_ok=True)
(OUT / "repeatability_check.json").unlink(missing_ok=True)
(OUT / "verification_report.json").unlink(missing_ok=True)
versions = {p: metadata.version(p) for p in ["numpy", "pandas", "scikit-learn", "matplotlib"]}
print(sys.version, sys.executable, versions, sep="\\n")
""")

step("3. Load and validate the dataset", """
Action: read the CSV and define the sensor groups explicitly. There are 19 columns: Date, Time, 16 sensor features and the occupancy target. Sixteen features do not mean sixteen physical sensors.
Temperature: S1-S4 Temp. Light: S1-S4 Light. Sound: S1-S4 Sound. CO2: S5_CO2 and S5_CO2_Slope together. PIR: S6_PIR and S7_PIR. Consult UCI documentation for units/calibration rather than assuming them from column names.
Use Date/Time only for ordering and evaluation, not prediction. Preserve source-row identifiers. Missing targets, invalid timestamps, unexpected labels or duplicate timestamps stop execution. Remove exact duplicate rows before splitting and count them.
Missing sensor readings are allowed and later median-imputed using training data only. Infinite values become missing. Random Forest does not require scaling.
Expected: 10,129 rows, four labels and no missing values for the supplied CSV. Checkpoint: read the audit before accepting any result.
""", """
GROUPS = {"Temperature": [f"S{i}_Temp" for i in range(1, 5)],
          "Light": [f"S{i}_Light" for i in range(1, 5)],
          "Sound": [f"S{i}_Sound" for i in range(1, 5)],
          "CO2": ["S5_CO2", "S5_CO2_Slope"], "PIR": ["S6_PIR", "S7_PIR"]}
FEATURES = sum(GROUPS.values(), [])
TARGET = "Room_Occupancy_Count"
raw = pd.read_csv(DATA)
assert set(raw.columns) == set(FEATURES + ["Date", "Time", TARGET])
assert raw[TARGET].notna().all() and raw[TARGET].isin(LABELS).all()
duplicates = int(raw.duplicated().sum())
df = raw.loc[~raw.duplicated()].copy()
df["source_row"] = df.index + 2  # CSV line number including the header.
df["timestamp"] = pd.to_datetime(df.Date + " " + df.Time, format="%Y/%m/%d %H:%M:%S", errors="raise")
assert df.timestamp.notna().all(), "Date and Time must not be missing."
assert not df.timestamp.duplicated().any(), "Investigate duplicate timestamps."
df = df.sort_values("timestamp").reset_index(drop=True)
df[FEATURES] = df[FEATURES].apply(pd.to_numeric, errors="raise").replace([np.inf, -np.inf], np.nan)
pir = df[["S6_PIR", "S7_PIR"]]
assert (pir.isna() | pir.isin([0, 1])).all().all()
audit = {"raw_shape": list(raw.shape), "duplicate_rows_removed": duplicates,
         "missing_sensor_values": int(df[FEATURES].isna().sum().sum()),
         "sha256": hashlib.sha256(DATA.read_bytes()).hexdigest(),
         "class_counts": df[TARGET].value_counts().sort_index().to_dict()}
print(json.dumps(audit, indent=2))
display(df.head())
daily = pd.crosstab(df.timestamp.dt.date, df[TARGET]).reindex(columns=LABELS, fill_value=0)
display(daily)
daily.to_csv(OUT / "daily_class_counts.csv")
""")

step("4. Explore the training data", """
Action: inspect full-data class balance and timing as an audit. Use the first recorded day for detailed sensor exploration: it is training data in every temporal evaluation. Do not use later feature patterns to tune the model.
Plot sensor distributions and Pearson correlations. Count readings beyond 1.5 times the interquartile range (IQR), but do not delete them automatically. Genuine occupied periods may look unusual relative to empty periods.
Expected: empty-room readings dominate (about 81%). Several dates contain only class 0, and January has no class 1. Checkpoint: explain why accuracy alone is insufficient and why test-class coverage matters.
""", """
def save_figure(name):
    plt.savefig(OUT / name, dpi=150, bbox_inches="tight")
    plt.show()
    plt.close()

fig, axes = plt.subplots(1, 2, figsize=(11, 3.5))
df[TARGET].value_counts().sort_index().plot.bar(ax=axes[0], color="#176b87")
axes[0].set(title="Full-data class audit", xlabel="Occupants", ylabel="Rows")
axes[1].scatter(df.timestamp, df[TARGET], s=2, color="#176b87")
axes[1].set(title="Occupancy over time", ylabel="Occupants")
fig.autofmt_xdate()
fig.tight_layout()
save_figure("class_timeline.png")
second_day = sorted(df.timestamp.dt.date.unique())[1]
eda_end = df.loc[df.timestamp.dt.date == second_day, "timestamp"].min() - pd.Timedelta(minutes=10)
eda = df.loc[(df.timestamp.dt.date == df.timestamp.dt.date.min()) & (df.timestamp < eda_end), FEATURES]
eda.hist(figsize=(12, 9), bins=25, color="#176b87")
plt.suptitle("Sensor distributions: first training day only")
plt.tight_layout()
save_figure("distributions.png")
fig, ax = plt.subplots(figsize=(9, 8))
im = ax.imshow(eda.corr(), vmin=-1, vmax=1, cmap="coolwarm")
ax.set_xticks(range(16), FEATURES, rotation=90)
ax.set_yticks(range(16), FEATURES)
ax.set_title("Pearson correlations: first training day")
fig.colorbar(im, ax=ax)
fig.tight_layout()
save_figure("correlations.png")
q1, q3 = eda.quantile(.25), eda.quantile(.75)
iqr = q3 - q1
flags = ((eda < q1 - 1.5 * iqr) | (eda > q3 + 1.5 * iqr)).sum()
display(flags.rename("IQR flags (not deleted)").to_frame())
flags.to_csv(OUT / "outlier_audit.csv")
""")

step("5. Define the controlled experiments", """
Action: compare All features with Without Temperature, Without Light, Without Sound, Without CO2, and Without PIR. A fresh pipeline is fitted for every experiment using identical split rows and model settings.
Use 100 trees, maximum depth 10, random seed 42, and max_features='sqrt', retaining the original settings. These are fixed instructional settings, not selected using test scores. Other key defaults are explicit below.
The sqrt rule is constant, but its effective candidate-feature count can change after removing columns. Ablation measures removal within this learning procedure, not a pure causal effect of a physical sensor.
Checkpoint: CO2 removal must include its slope. Date, Time, source_row, timestamp and the target must never enter X.
""", """
CONFIGS = {"All features": FEATURES}
CONFIGS.update({f"Without {g}": [f for f in FEATURES if f not in cols] for g, cols in GROUPS.items()})
PARAMS = dict(n_estimators=100, max_depth=10, random_state=SEED, n_jobs=-1,
              criterion="gini", max_features="sqrt", bootstrap=True,
              min_samples_split=2, min_samples_leaf=1, class_weight=None)
display(pd.DataFrame([{"configuration": name, "n_features": len(cols),
                       "removed": "No group removed" if name == "All features" else name[8:]}
                      for name, cols in CONFIGS.items()]))
assert len(FEATURES) == len(set(FEATURES)) == 16
""")

step("6. Create and audit the splits", """
Action: preserve the stratified random 80:20 split (seed 42) as a historical comparison, not evidence of future-session generalization. Neighboring readings can appear on both sides and are highly dependent.
Temporal holdout: train on December, test on January. This uses the natural multi-day recording gap. January has only classes 0, 2 and 3, so class-1 generalization cannot be assessed. No test scores are used to tune models.
Robustness: train on expanding past history and test each next recorded day, with a 10-minute embargo before that day's first test reading. This illustrative gap reduces immediate adjacency; it does not prove independence.
No fold is selected by performance. Separate occupied-day and empty-only-day summaries. Folds share training history and are not independent replications. The January holdout overlaps the final two daily tests, so these are alternative summaries, not independent confirmation.
Checkpoint: assert disjoint rows, temporal ordering and class coverage. Save exact source-row assignments. The cutoff and embargo are study design choices specific to this dataset, not optimized settings.
""", """
indices = np.arange(len(df))
tr, te = train_test_split(indices, test_size=.2, stratify=df[TARGET], random_state=SEED)
SPLITS = {"random_80_20": (tr, te)}
cutoff = pd.Timestamp("2018-01-01")
SPLITS["temporal_holdout"] = (indices[df.timestamp < cutoff], indices[df.timestamp >= cutoff])
dates = sorted(df.timestamp.dt.date.unique())
for day in dates[1:]:
    te = indices[df.timestamp.dt.date == day]
    tr = indices[df.timestamp < df.timestamp.iloc[te].min() - pd.Timedelta(minutes=10)]
    SPLITS[f"daily_{day}"] = (tr, te)
split_rows, assignment_rows = [], []
for name, (tr, te) in SPLITS.items():
    assert len(tr) and len(te) and not set(tr).intersection(te)
    assert set(df[TARGET].iloc[te]).issubset(set(df[TARGET].iloc[tr]))
    if name != "random_80_20":
        assert df.timestamp.iloc[tr].max() < df.timestamp.iloc[te].min()
    split_rows.append({"split": name, "train_rows": len(tr), "test_rows": len(te),
                       "train_end": str(df.timestamp.iloc[tr].max()),
                       "test_start": str(df.timestamp.iloc[te].min()),
                       "test_classes": ",".join(map(str, sorted(df[TARGET].iloc[te].unique()))),
                       **{f"train_class_{c}": int((df[TARGET].iloc[tr] == c).sum()) for c in LABELS},
                       **{f"test_class_{c}": int((df[TARGET].iloc[te] == c).sum()) for c in LABELS}})
    for role, rows in [("train", tr), ("test", te)]:
        assignment_rows.extend({"split": name, "role": role, "source_row": int(df.source_row.iloc[i])} for i in rows)
split_audit = pd.DataFrame(split_rows)
display(split_audit)
split_audit.to_csv(OUT / "split_audit.csv", index=False)
pd.DataFrame(assignment_rows).to_csv(OUT / "split_assignments.csv", index=False)
""")

step("7. Train and evaluate all configurations", """
Action: run 48 small forests, six configurations across eight splits. Runtime depends on your machine. The median imputer is fitted inside each training pipeline; no scaling or oversampling is required.
Accuracy is the fraction correct. Precision asks how often a predicted class is correct; recall asks how often an actual class is found. F1 balances precision and recall. Weighted averages weight by class support and are dominated by class 0 here. Weighted recall equals accuracy in this single-label task.
Macro-F1 is the arithmetic mean of per-class F1 values. Use the classes present in each test split for the primary within-split ablation comparison. Also export fixed-four-class macro-F1, assigning zero to absent classes, for transparency. The report label macro avg fixed4 explicitly identifies that alternative average. Never compare macro scores across splits without considering class coverage. A zero-support class is untested, not a measured failure.
Export per-class reports, confusion matrices and exact predictions for every experiment. Checkpoint: each confusion matrix sums to the test-row count. Do not retrain just to print a report.
""", """
rows, prediction_rows, reports, matrices = [], [], {}, {}
importance_tables = []
for split, (tr, te) in SPLITS.items():
    truth = df[TARGET].iloc[te]
    observed = sorted(truth.unique())
    reports[split], matrices[split] = {}, {}
    for name, cols in CONFIGS.items():
        assert df.iloc[tr][cols].notna().any().all(), "All-missing training feature."
        model = make_pipeline(SimpleImputer(strategy="median"), RandomForestClassifier(**PARAMS))
        model.fit(df.iloc[tr][cols], df[TARGET].iloc[tr])
        pred = model.predict(df.iloc[te][cols])
        pw, rw, fw, _ = precision_recall_fscore_support(truth, pred, average="weighted", zero_division=0)
        pm, rm, fm, _ = precision_recall_fscore_support(truth, pred, labels=observed, average="macro", zero_division=0)
        f4 = precision_recall_fscore_support(truth, pred, labels=LABELS, average="macro", zero_division=0)[2]
        rows.append({"split": split, "configuration": name, "features": len(cols),
                     "accuracy": accuracy_score(truth, pred), "precision_weighted": pw,
                     "recall_weighted": rw, "f1_weighted": fw,
                     "precision_macro_observed": pm, "recall_macro_observed": rm,
                     "f1_macro_observed": fm, "f1_macro_fixed4": f4,
                     "errors": int((truth.to_numpy() != pred).sum())})
        report = classification_report(truth, pred, labels=LABELS, output_dict=True, zero_division=0)
        report["macro avg fixed4"] = report.pop("macro avg")
        report["macro avg observed"] = {"precision": pm, "recall": rm, "f1-score": fm, "support": len(te)}
        reports[split][name] = report
        cm = confusion_matrix(truth, pred, labels=LABELS)
        assert cm.sum() == len(te)
        matrices[split][name] = cm.tolist()
        prediction_rows.extend({"split": split, "configuration": name,
                                "source_row": int(df.source_row.iloc[i]), "actual": int(y), "predicted": int(p)}
                               for i, y, p in zip(te, truth, pred))
        if name == "All features":
            importance_tables.extend({"split": split, "feature": f, "importance": float(v),
                                      "group": next(g for g, fs in GROUPS.items() if f in fs)}
                                     for f, v in zip(cols, model[-1].feature_importances_))
    print("Completed", split)
metrics = pd.DataFrame(rows)
base = metrics[metrics.configuration == "All features"].set_index("split")
metrics["accuracy_drop_pp"] = 100 * (metrics.split.map(base.accuracy) - metrics.accuracy)
metrics["macro_f1_drop_pp"] = 100 * (metrics.split.map(base.f1_macro_observed) - metrics.f1_macro_observed)
metrics.to_csv(OUT / "metrics.csv", index=False)
pd.DataFrame(prediction_rows).to_csv(OUT / "predictions.csv", index=False)
(OUT / "classification_reports.json").write_text(json.dumps(reports, indent=2))
(OUT / "confusion_matrices.json").write_text(json.dumps(matrices, indent=2))
display(metrics[metrics.split.isin(["random_80_20", "temporal_holdout"])])
""")

step("8. Inspect errors and ablation changes", """
Action: inspect all six confusion matrices for the random comparison and temporal holdout. Rows are true labels, columns predictions. Off-diagonal entries are mistakes; an empty true-label row means there were no examples, not perfect recall.
Positive drop means removal worsens performance; negative means improvement. Drops use percentage points: 99.70% minus 99.16% is about 0.54 percentage points, not a relative decrease of 0.54%.
Checkpoint: examine occupied-class errors rather than just correct-empty-room counts. The baseline bar is colored by name, regardless of sorting. Full per-class reports for all configurations are saved in classification_reports.json.
""", """
for split in ["random_80_20", "temporal_holdout"]:
    fig, axes = plt.subplots(2, 3, figsize=(12, 7))
    for ax, name in zip(axes.flat, CONFIGS):
        ConfusionMatrixDisplay(np.array(matrices[split][name]), display_labels=LABELS).plot(ax=ax, colorbar=False, cmap="Blues")
        ax.set_title(name)
    fig.suptitle(split + ": rows=true, columns=predicted")
    fig.tight_layout()
    save_figure(f"confusion_{split}.png")
    subset = metrics[metrics.split == split].sort_values("f1_macro_observed")
    fig, ax = plt.subplots(figsize=(9, 4))
    ax.barh(subset.configuration, subset.f1_macro_observed,
            color=["#bf622f" if n == "All features" else "#176b87" for n in subset.configuration])
    ax.set(xlim=(0, 1.08), xlabel="Macro-F1 (observed classes)", title=split + ": orange = baseline")
    for i, value in enumerate(subset.f1_macro_observed):
        ax.text(value + .01, i, f"{value:.4f}", va="center")
    fig.tight_layout()
    save_figure(f"comparison_{split}.png")
    for name in CONFIGS:
        print(split, name, "per-class report: support=0 means untested")
        display(pd.DataFrame(reports[split][name]).T)
""")

step("9. Analyze feature and group importance", """
Action: rank baseline impurity-based feature importance and sum within groups for both principal splits. These describe the trained forest, not causality or standalone sensor usefulness.
Impurity importance can favor features with many distinct values. Correlated inputs may substitute for one another. A feature can be important in the full model yet have a small ablation penalty because retraining lets other inputs replace it. Group sums also depend on group size.
Checkpoint: use ablation as main conditional-contribution evidence and importance as support. Do not declare Sound noisy merely because removing it improves a score.
""", """
importance = pd.DataFrame(importance_tables)
group_importance = importance.groupby(["split", "group"], as_index=False).importance.sum()
assert np.allclose(group_importance.groupby("split").importance.sum(), 1)
importance.to_csv(OUT / "feature_importance.csv", index=False)
group_importance.to_csv(OUT / "group_importance.csv", index=False)
for split in ["random_80_20", "temporal_holdout"]:
    display(importance[importance.split == split].sort_values("importance", ascending=False))
    fig, axes = plt.subplots(1, 2, figsize=(12, 6))
    for ax, frame, label in [(axes[0], importance, "feature"), (axes[1], group_importance, "group")]:
        part = frame[frame.split == split].sort_values("importance")
        ax.barh(part[label], part.importance, color="#176b87")
        ax.set_xlabel("Impurity importance")
    fig.suptitle(split + ": baseline supporting evidence")
    fig.tight_layout()
    save_figure(f"importance_{split}.png")
""")

step("10. Assess robustness across days", """
Action: summarize paired ablation drops within expanding-history daily folds. Report mean, sample standard deviation, minimum and maximum, separating occupied-present from empty-only days. Days are weighted equally, unlike pooled row-level scoring.
Only two daily test folds include occupied readings; four are empty-only. Test classes differ and training history overlaps. The spread is descriptive variability, not a confidence interval or statistical significance test. Repeating random seeds alone would not resolve temporal leakage.
Checkpoint: does the ranking change between random and temporal evaluations? Do not claim a universal most/least essential group when the evidence differs. Future work needs more independent occupied sessions and verification that the supplied CO2 slope is computed causally for real-time use.
""", """
daily_metrics = metrics[metrics.split.str.startswith("daily_")].copy()
test_classes = split_audit.set_index("split").test_classes
daily_metrics["day_type"] = daily_metrics.split.map(test_classes).map(lambda s: "empty-only" if s == "0" else "occupied-present")
robustness = (daily_metrics.groupby(["day_type", "configuration"])["macro_f1_drop_pp"]
              .agg(["count", "mean", "std", "min", "max"]).reset_index())
robustness.to_csv(OUT / "robustness.csv", index=False)
display(robustness)
display(daily_metrics[["split", "configuration", "accuracy", "f1_macro_observed", "macro_f1_drop_pp", "day_type"]])
""")

step("11. Write evidence-based conclusions", """
Action: report the exact split, test-class support, baseline scores and the largest/smallest observed ablation drops. Smallest does not prove dispensability. Zero or negative drops do not establish lack of information.
The initial random split classified 2,020 of 2,026 correctly. Removing Sound improved two predictions, not evidence of noise reduction. Compare the regenerated table instead of assuming time-separated evaluation preserves that ranking.
Limitations: one room, few recording days, temporal dependence, severe imbalance, absent class 1 in January, unverified real-time availability of CO2 slope and no external validation. This is offline classification, not a certified occupancy or safety system.
Checkpoint: address all three objectives while stating where evidence is insufficient. More occupied-session data is a recommendation, not an already completed experiment.
""", """
conclusions = []
for split in ["random_80_20", "temporal_holdout"]:
    b = base.loc[split]
    ab = metrics[(metrics.split == split) & (metrics.configuration != "All features")]
    largest = ab.loc[ab.macro_f1_drop_pp.idxmax()]
    smallest = ab.loc[ab.macro_f1_drop_pp.idxmin()]
    text = (f"{split}: baseline accuracy={b.accuracy:.4%}, macro-F1 observed={b.f1_macro_observed:.4f}; "
            f"largest observed macro-F1 drop: {largest.configuration} ({largest.macro_f1_drop_pp:+.3f} pp); "
            f"smallest: {smallest.configuration} ({smallest.macro_f1_drop_pp:+.3f} pp). "
            "These are conditional effects, not universal sensor-necessity rankings.")
    conclusions.append(text)
    print(text)
(OUT / "conclusions.json").write_text(json.dumps(conclusions, indent=2))
""")

step("12. Reproduce and submit", """
Action: choose Restart Kernel and Run All Cells in Jupyter. Wait for completion and inspect outputs and warnings. The notebook overwrites its named outputs in results/. Never edit scores manually.
To rebuild the notebook and guides, run python build_materials.py. To regenerate guides after running the notebook, run python build_materials.py --documents-only. To execute twice and compare all numeric experiment outputs, run python build_materials.py --verify.
Keep requirements.txt, build_materials.py, notebook, CSV and results/ together. The DOCX opens in Word or imports into Google Docs; this does not create an online Google Doc. Older root-level reports and result CSVs are superseded by the current guide and results/.
Submission checklist: executed notebook; DOCX/PDF guide; checksum and versions; split audit/assignments; all six configurations per split; per-class reports/matrices; individual/group importance; robustness summaries; qualified conclusions.
Student questions: Why can random splitting inflate sensor-data performance? Why does weighted recall equal accuracy? How can a feature be important but have little ablation impact? Why is January class-1 recall untested? What further data supports deployment?
Troubleshooting: FileNotFoundError means working directory or CSV path is wrong. ModuleNotFoundError often means packages and kernel use different environments. Invalid labels/timestamps require investigation, not disabling assertions. PDF generation uses ReportLab and needs no LaTeX, browser or Pango.
Checkpoint: manifest and assertions must complete. Identical numeric results are expected in the same tested environment; different library versions can change results. The pinned file records direct dependencies, not every transitive package.
References: supplied Intro-to-AI-Proposal.docx; UCI Room Occupancy Estimation dataset, https://archive.ics.uci.edu/dataset/864/room+occupancy+estimation ; scikit-learn RandomForestClassifier, classification metrics and common pitfalls documentation, https://scikit-learn.org/stable/ .
""", """
assert len(metrics) == len(CONFIGS) * len(SPLITS) == 48
assert np.allclose(metrics.accuracy, metrics.recall_weighted)
assert (metrics.loc[metrics.configuration == "All features", "macro_f1_drop_pp"] == 0).all()
manifest = {"python": sys.version, "versions": versions, "dataset": audit,
            "builder_sha256": hashlib.sha256(Path("build_materials.py").read_bytes()).hexdigest(),
            "random_seed": SEED, "forest_parameters": PARAMS,
            "temporal_cutoff": str(cutoff), "daily_embargo_minutes": 10,
            "configurations": CONFIGS, "split_count": len(SPLITS),
            "author": "Prepared by Engr. Jamie Eduardo Rosal, MSCpE"}
(OUT / "manifest.json").write_text(json.dumps(manifest, indent=2))
print("Completed: 48 experiments; outputs and manifest saved in results/.")
""")


def build_notebook():
    nb = nbformat.v4.new_notebook()
    nb.metadata.kernelspec = {"name": "python3", "language": "python", "display_name": "Python 3"}
    nb.cells = [nbformat.v4.new_markdown_cell(f"# {TITLE}\n\n{AUTHOR}\n\nRun in order. These 12 steps match the DOCX/PDF guide and supersede the initial random-split-only report.")]
    for title, explanation, code in STEPS:
        nb.cells.extend([nbformat.v4.new_markdown_cell("## " + title + "\n\n" + explanation.replace("\n", "\n\n")), nbformat.v4.new_code_cell(code)])
    # Match the kernel to the builder's interpreter, not an unrelated default Python.
    from jupyter_client import KernelManager
    km = KernelManager(kernel_name="python3")
    km.kernel_spec.argv = [sys.executable, "-m", "ipykernel_launcher", "-f", "{connection_file}"]
    try:
        NotebookClient(nb, timeout=600, km=km, resources={"metadata": {"path": str(ROOT)}}).execute()
    finally:
        if km.has_kernel:
            km.shutdown_kernel(now=True)
    nbformat.validate(nb)
    assert all(c.execution_count is not None for c in nb.cells if c.cell_type == "code")
    assert not any(o.output_type == "error" for c in nb.cells if c.cell_type == "code" for o in c.outputs)
    nbformat.write(nb, ROOT / "Room_Occupancy_Estimation.ipynb")
    from nbconvert import HTMLExporter
    html, _ = HTMLExporter().from_notebook_node(nb)
    (ROOT / "Room_Occupancy_Estimation.html").write_text(html)


def documents():
    out = ROOT / "results"
    metrics = pd.read_csv(out / "metrics.csv")
    manifest = json.loads((out / "manifest.json").read_text())
    assert manifest["dataset"]["sha256"] == hashlib.sha256((ROOT / "Occupancy_Estimation.csv").read_bytes()).hexdigest()
    assert manifest["builder_sha256"] == hashlib.sha256(Path(__file__).read_bytes()).hexdigest(), "Builder changed: run a full build, not documents-only."
    doc = Document()
    doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(10)
    doc.sections[0].header.paragraphs[0].text = "INTRODUCTION TO AI | STUDENT LAB GUIDE"
    doc.sections[0].footer.paragraphs[0].text = AUTHOR
    styles = getSampleStyleSheet()
    styles["BodyText"].leading = 14
    story = []

    def paragraph(text, heading=0):
        if heading:
            doc.add_heading(text, level=heading)
        else:
            doc.add_paragraph(text)
        story.append(Paragraph(escape(text), styles[f"Heading{heading}"] if heading else styles["BodyText"]))
        space = Spacer(1, 6)
        space.keepWithNext = bool(heading)
        story.append(space)

    def table(frame):
        data = [list(frame.columns)] + frame.astype(str).values.tolist()
        t = doc.add_table(rows=1, cols=len(frame.columns))
        t.style = "Light Shading Accent 1"
        from docx.oxml import OxmlElement
        repeat_header = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat_header)
        for cell, value in zip(t.rows[0].cells, data[0]):
            cell.text = value
        for row in data[1:]:
            for cell, value in zip(t.add_row().cells, row):
                cell.text = value
        formatted = [[Paragraph(escape(v), styles["BodyText"]) for v in row] for row in data]
        pt = Table(formatted, colWidths=[468 / len(data[0])] * len(data[0]), repeatRows=1)
        pt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d9eaf0")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("GRID", (0, 0), (-1, -1), .4, colors.lightgrey),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        leading = []
        if len(story) >= 2 and isinstance(story[-2], Paragraph) and story[-2].style.name.startswith("Heading"):
            leading = story[-2:]
            del story[-2:]
        story.extend([KeepTogether(leading + [pt]), Spacer(1, 12)])

    def image(filename, caption):
        from PIL import Image as PILImage
        path = out / filename
        with PILImage.open(path) as pic:
            w, h = pic.size
        width = min(468, 510 * w / h)
        doc.add_picture(str(path), width=Inches(width / 72))
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
        doc.add_paragraph(caption)
        story.append(KeepTogether([Image(str(path), width=width, height=width * h / w),
                                   Paragraph(escape(caption), styles["BodyText"]), Spacer(1, 12)]))

    paragraph(TITLE, 1)
    paragraph(AUTHOR)
    paragraph("A repeatable laboratory workflow aligned with the supplied AI proposal. Revised edition: time-aware evaluation, controlled sensor ablation and cautious interpretation.")
    paragraph("How to use this guide", 2)
    paragraph("Work through steps 1-12 in order. Read the action and checkpoint, run the matching notebook cell, then inspect its output. The guide provides instructions; the notebook contains executable code. The tables below are generated from notebook outputs, not typed manually.")
    for title, _, _ in STEPS:
        paragraph(title)
    for title, explanation, _ in STEPS:
        doc.add_page_break()
        story.append(PageBreak())
        paragraph(title, 1)
        for line in explanation.splitlines():
            paragraph(line)
        if title.startswith("2."):
            paragraph("Commands after activating the environment", 2)
            for command in ["python -m pip install -r requirements.txt", 'python -m ipykernel install --user --name occupancy-guide --display-name "Occupancy Guide"', "python -m jupyterlab"]:
                paragraph(command)
        if title.startswith("3."):
            table(pd.read_csv(out / "daily_class_counts.csv"))
        if title.startswith("4."):
            image("class_timeline.png", "Class and timing audit. Recording gaps are not interpolated.")
            image("distributions.png", "First-day training distributions; unusual readings are retained.")
            image("correlations.png", "Correlation is descriptive, not causal.")
        if title.startswith("6."):
            audit = pd.read_csv(out / "split_audit.csv")
            table(audit[["split", "train_rows", "test_rows", "test_classes"]])
        if title.startswith("7."):
            for split in ["random_80_20", "temporal_holdout"]:
                paragraph(split, 2)
                part = metrics[metrics.split == split]
                table(pd.DataFrame({"Configuration": part.configuration, "Accuracy": part.accuracy.map(lambda v: f"{v:.4f}"),
                                    "Weighted F1": part.f1_weighted.map(lambda v: f"{v:.4f}"),
                                    "Macro F1 observed": part.f1_macro_observed.map(lambda v: f"{v:.4f}"),
                                    "Macro drop pp": part.macro_f1_drop_pp.map(lambda v: f"{v:+.3f}")}))
        if title.startswith("8."):
            for split in ["random_80_20", "temporal_holdout"]:
                image(f"confusion_{split}.png", split + ": all six configurations; zero-support rows are untested.")
                image(f"comparison_{split}.png", split + ": orange highlights the baseline.")
                paragraph(split + " baseline per-class scores", 2)
                report = json.loads((out / "classification_reports.json").read_text())[split]["All features"]
                table(pd.DataFrame([{"Class": str(c), "Precision": f"{report[str(c)]['precision']:.4f}",
                                     "Recall": f"{report[str(c)]['recall']:.4f}",
                                     "F1": f"{report[str(c)]['f1-score']:.4f}",
                                     "Support": str(int(report[str(c)]['support']))} for c in range(4)]))
        if title.startswith("9."):
            for split in ["random_80_20", "temporal_holdout"]:
                image(f"importance_{split}.png", split + ": individual-feature and group-level supporting evidence.")
        if title.startswith("10."):
            robust = pd.read_csv(out / "robustness.csv")
            for kind in robust.day_type.unique():
                paragraph(kind + " daily folds", 2)
                r = robust[robust.day_type == kind]
                table(pd.DataFrame({"Configuration": r.configuration, "Folds": r["count"],
                                    "Mean drop pp": r["mean"].map(lambda v: f"{v:+.3f}"),
                                    "SD pp": r["std"].map(lambda v: f"{v:.3f}"),
                                    "Range pp": [f"{a:+.3f} to {b:+.3f}" for a, b in zip(r["min"], r["max"])]}))
        if title.startswith("11."):
            paragraph("Computed findings", 2)
            for text in json.loads((out / "conclusions.json").read_text()):
                paragraph(text)
        if title.startswith("12."):
            paragraph("Dataset SHA-256", 2)
            checksum = manifest["dataset"]["sha256"]
            paragraph(checksum[:32] + " " + checksum[32:] + " (join without the space)")
            paragraph("Tested analysis versions: " + json.dumps(manifest["versions"]))
    doc.save(ROOT / "Room_Occupancy_Student_Guide.docx")

    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawString(72, 35, "Room Occupancy Estimation | Student Guide")
        canvas.drawRightString(540, 35, str(document.page))

    SimpleDocTemplate(str(ROOT / "Room_Occupancy_Student_Guide.pdf"), rightMargin=72, leftMargin=72,
                      topMargin=54, bottomMargin=54, pagesize=(612, 792), title=TITLE,
                      author=AUTHOR).build(story, onFirstPage=footer, onLaterPages=footer)
    from pypdf import PdfReader
    pdf = PdfReader(ROOT / "Room_Occupancy_Student_Guide.pdf")
    text = "\n".join(p.extract_text() for p in pdf.pages)
    for title, _, _ in STEPS:
        assert title in text, title
    assert AUTHOR in text and "nan" not in text.lower().split()
    print(f"DOCX/PDF generated and text-checked: {len(pdf.pages)} PDF pages.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--documents-only", action="store_true")
    parser.add_argument("--verify", action="store_true")
    args = parser.parse_args()
    if not args.documents_only:
        build_notebook()
        if args.verify:
            names = ["metrics.csv", "predictions.csv", "feature_importance.csv", "group_importance.csv",
                     "split_assignments.csv", "split_audit.csv", "robustness.csv",
                     "classification_reports.json", "confusion_matrices.json"]
            hashes = {n: hashlib.sha256((ROOT / "results" / n).read_bytes()).hexdigest() for n in names}
            build_notebook()
            assert hashes == {n: hashlib.sha256((ROOT / "results" / n).read_bytes()).hexdigest() for n in names}
            (ROOT / "results" / "repeatability_check.json").write_text(json.dumps({"status": "PASS", "matching_sha256": hashes}, indent=2))
            print("PASS: identical metrics, predictions, importances, splits and reports across two executions.")
    documents()


if __name__ == "__main__":
    main()
