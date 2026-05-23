from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
DATA_FILE = ROOT / "AITHOS-SURVEY-RAW-DATA.xlsx"
REPORT_FILE = ROOT / "AITHOS_Statistical_Analysis_Report.docx"
NOTEBOOK_FILE = ROOT / "AITHOS_Statistical_Computations_Validation.ipynb"
FIGURES_DIR = ROOT / "figures"

ALPHA = 0.05

CONSTRUCTS = {
    "Expertise": {
        "sheet": "PART II.",
        "cols": [1, 2, 3, 4, 5],
        "domain": "AI Credibility",
        "items": [
            "AI provides accurate information regarding coffee shop trends.",
            "AI-generated content demonstrates a high level of knowledge about the industry.",
            "The suggestions made by the AI appear professional.",
            "AI-generated suggestions are realistic for business applications.",
            "The AI quickly provides relevant solutions for coffee shop promotion.",
        ],
    },
    "Trustworthiness": {
        "sheet": "PART II.",
        "cols": [7, 8, 9, 10, 11],
        "domain": "AI Credibility",
        "items": [
            "AI-generated content is perceived as trustworthy.",
            "AI-generated content can provide unbiased marketing suggestions.",
            "AI-generated content aligns with the brand's values.",
            "AI-generated content can be trusted for customer communication.",
            "AI-generated content provides dependable communication support.",
        ],
    },
    "Algorithmic Transparency": {
        "sheet": "PART III.",
        "cols": [1, 2, 3, 4, 5],
        "domain": "Algorithmic Transparency",
        "items": [
            "It is clear how the AI decides what content to generate.",
            "The respondent understands why AI produces specific results from prompts.",
            "The way the AI functions is not unfamiliar.",
            "The way AI creates suggestions is easy to follow.",
            "The respondent knows the limitations and weak spots of the AI tool used.",
        ],
    },
    "Perceived Usefulness": {
        "sheet": "PART IV.",
        "cols": [1, 2, 3, 4, 5],
        "domain": "Technology Acceptance",
        "items": [
            "Using AI enhances productivity in managing the coffee shop.",
            "AI tools help reach marketing goals more effectively.",
            "AI tools are useful in daily business operations.",
            "AI tools assist in identifying new business opportunities.",
            "Relying on AI helps handle unexpected business challenges more effectively.",
        ],
    },
    "Perceived Ease of Use": {
        "sheet": "PART IV.",
        "cols": [7, 8, 9, 10, 11],
        "domain": "Technology Acceptance",
        "items": [
            "Interacting with AI tools does not require much mental effort.",
            "It is easy to get AI to do what the respondent wants it to do.",
            "Learning to operate AI tools is easy.",
            "AI tools can be accessed without complicated setup.",
            "The respondent does not need help to figure out how to use AI tools properly.",
        ],
    },
    "Quality of Content": {
        "sheet": "PART V.",
        "cols": [1, 2, 3, 4, 5],
        "domain": "Communication Strategies",
        "items": [
            "AI-generated content is clear to understand and relevant.",
            "AI-generated content meets quality standards before posting.",
            "AI-generated content provides reliable information to customers.",
            "AI organizes text so online readers can scan it easily.",
            "AI quickly gives content that fits current trends or seasons.",
        ],
    },
    "Message Consistency": {
        "sheet": "PART V.",
        "cols": [7, 8, 9, 10, 11],
        "domain": "Communication Strategies",
        "items": [
            "AI-assisted messages align with business values.",
            "Customer communication remains consistent when using AI.",
            "AI helps deliver messages across different platforms.",
            "Customers are not confused by information AI helps write.",
            "AI content sounds like it was written by a real staff member.",
        ],
    },
    "Audience Engagement": {
        "sheet": "PART V.",
        "cols": [14, 15, 16, 17, 18],
        "domain": "Communication Strategies",
        "items": [
            "AI helps increase customer engagement.",
            "AI helps respond quickly to customer inquiries.",
            "AI encourages customer feedback and interaction.",
            "AI-written captions or promos make customers want to visit or buy.",
            "AI writes captions using language and trends local coffee lovers understand.",
        ],
    },
}


def verbal_interpretation(mean: float) -> str:
    if mean >= 3.26:
        return "Strongly Agree / Very High"
    if mean >= 2.51:
        return "Agree / High"
    if mean >= 1.76:
        return "Disagree / Low"
    return "Strongly Disagree / Very Low"


def alpha_interpretation(alpha: float) -> str:
    if pd.isna(alpha):
        return "Not computable"
    if alpha >= 0.90:
        return "Excellent"
    if alpha >= 0.80:
        return "Good"
    if alpha >= 0.70:
        return "Acceptable"
    if alpha >= 0.60:
        return "Questionable"
    return "Low"


def correlation_interpretation(r: float) -> str:
    ar = abs(r)
    if ar >= 0.90:
        strength = "Very strong"
    elif ar >= 0.70:
        strength = "Strong"
    elif ar >= 0.50:
        strength = "Moderate"
    elif ar >= 0.30:
        strength = "Weak"
    elif ar >= 0.10:
        strength = "Very weak"
    else:
        strength = "Negligible"
    direction = "positive" if r >= 0 else "negative"
    return f"{strength} {direction}"


def p_value_from_r(r: float, n: int) -> float:
    if n <= 3 or not np.isfinite(r):
        return float("nan")
    r = min(max(float(r), -0.999999999), 0.999999999)
    z = math.atanh(r) * math.sqrt(n - 3)
    return math.erfc(abs(z) / math.sqrt(2.0))


def cronbach_alpha(df: pd.DataFrame) -> float:
    clean = df.dropna()
    k = clean.shape[1]
    if k < 2 or clean.shape[0] < 2:
        return float("nan")
    item_variances = clean.var(axis=0, ddof=1)
    total_variance = clean.sum(axis=1).var(ddof=1)
    if total_variance == 0:
        return float("nan")
    return float((k / (k - 1)) * (1 - item_variances.sum() / total_variance))


def pearson_pair(x: pd.Series, y: pd.Series) -> dict:
    pair = pd.concat([x, y], axis=1).dropna()
    n = len(pair)
    r = float(pair.iloc[:, 0].corr(pair.iloc[:, 1])) if n >= 2 else float("nan")
    p = p_value_from_r(r, n)
    return {
        "n": n,
        "r": r,
        "p_value": p,
        "decision": "Significant" if p < ALPHA else "Not significant",
        "interpretation": correlation_interpretation(r),
    }


def clean_constructs() -> tuple[pd.DataFrame, dict[str, pd.DataFrame]]:
    raw = pd.read_excel(DATA_FILE, sheet_name=None, header=None)
    construct_items = {}
    respondent_ids = None

    for name, spec in CONSTRUCTS.items():
        sheet = raw[spec["sheet"]]
        block = sheet.iloc[9:, [0] + spec["cols"]].copy()
        block.columns = ["Respondent"] + [f"{name} {i}" for i in range(1, 6)]
        block = block.dropna(how="all")
        block["Respondent"] = pd.to_numeric(block["Respondent"], errors="coerce").astype("Int64")
        for col in block.columns[1:]:
            block[col] = pd.to_numeric(block[col], errors="coerce")
        block = block.dropna(subset=["Respondent"]).set_index("Respondent").sort_index()
        construct_items[name] = block
        if respondent_ids is None:
            respondent_ids = block.index
        elif not respondent_ids.equals(block.index):
            raise ValueError(f"Respondent IDs do not align for {name}")

    scored = pd.DataFrame(index=respondent_ids)
    for name, block in construct_items.items():
        scored[name] = block.mean(axis=1)

    scored["Overall Credibility"] = scored[["Expertise", "Trustworthiness"]].mean(axis=1)
    scored["Overall Technology Acceptance"] = scored[["Perceived Usefulness", "Perceived Ease of Use"]].mean(axis=1)
    scored["Overall Communication Strategy"] = scored[
        ["Quality of Content", "Message Consistency", "Audience Engagement"]
    ].mean(axis=1)

    return scored, construct_items


def validate_data(scored: pd.DataFrame, construct_items: dict[str, pd.DataFrame]) -> list[str]:
    messages = []
    messages.append(f"Respondent count: {len(scored)}")
    messages.append(f"Respondent IDs aligned: {scored.index.min()} to {scored.index.max()}")
    all_values = pd.concat(construct_items.values(), axis=1)
    bad = all_values[(all_values < 1) | (all_values > 4)].count().sum()
    missing = all_values.isna().sum().sum()
    messages.append(f"Out-of-range Likert values: {int(bad)}")
    messages.append(f"Missing Likert values: {int(missing)}")
    for name, block in construct_items.items():
        messages.append(f"{name}: {block.shape[1]} items, {block.shape[0]} respondents")
    return messages


def summarize_constructs(scored: pd.DataFrame) -> pd.DataFrame:
    rows = []
    ordered = [
        "Expertise",
        "Trustworthiness",
        "Overall Credibility",
        "Algorithmic Transparency",
        "Perceived Usefulness",
        "Perceived Ease of Use",
        "Overall Technology Acceptance",
        "Quality of Content",
        "Message Consistency",
        "Audience Engagement",
        "Overall Communication Strategy",
    ]
    for name in ordered:
        series = scored[name]
        rows.append(
            {
                "Construct": name,
                "Mean": series.mean(),
                "SD": series.std(ddof=1),
                "Min": series.min(),
                "Max": series.max(),
                "Interpretation": verbal_interpretation(series.mean()),
            }
        )
    return pd.DataFrame(rows)


def summarize_items(construct_items: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows = []
    for construct, block in construct_items.items():
        for idx, col in enumerate(block.columns, start=1):
            rows.append(
                {
                    "Construct": construct,
                    "Item": idx,
                    "Statement": CONSTRUCTS[construct]["items"][idx - 1],
                    "Mean": block[col].mean(),
                    "SD": block[col].std(ddof=1),
                    "Interpretation": verbal_interpretation(block[col].mean()),
                }
            )
    return pd.DataFrame(rows)


def summarize_reliability(construct_items: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows = []
    for construct, block in construct_items.items():
        alpha = cronbach_alpha(block)
        rows.append(
            {
                "Scale": construct,
                "Items": block.shape[1],
                "Cronbach Alpha": alpha,
                "Interpretation": alpha_interpretation(alpha),
            }
        )

    combined_scales = {
        "Overall AI Credibility": ["Expertise", "Trustworthiness"],
        "Overall Communication Strategy": ["Quality of Content", "Message Consistency", "Audience Engagement"],
        "Overall Technology Acceptance": ["Perceived Usefulness", "Perceived Ease of Use"],
    }
    for label, constructs in combined_scales.items():
        block = pd.concat([construct_items[c] for c in constructs], axis=1)
        alpha = cronbach_alpha(block)
        rows.append(
            {
                "Scale": label,
                "Items": block.shape[1],
                "Cronbach Alpha": alpha,
                "Interpretation": alpha_interpretation(alpha),
            }
        )
    return pd.DataFrame(rows)


def summarize_correlations(scored: pd.DataFrame) -> pd.DataFrame:
    predictors = ["Expertise", "Trustworthiness", "Overall Credibility"]
    outcomes = ["Quality of Content", "Message Consistency", "Audience Engagement", "Overall Communication Strategy"]
    rows = []
    for predictor in predictors:
        for outcome in outcomes:
            result = pearson_pair(scored[predictor], scored[outcome])
            rows.append(
                {
                    "Predictor": predictor,
                    "Outcome": outcome,
                    "n": result["n"],
                    "r": result["r"],
                    "p-value": result["p_value"],
                    "Decision": result["decision"],
                    "Interpretation": result["interpretation"],
                }
            )
    return pd.DataFrame(rows)


def fmt(value: float, digits: int = 3) -> str:
    if pd.isna(value):
        return "NA"
    return f"{value:.{digits}f}"


def load_font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Helvetica.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
    ]
    for path in candidates:
        if path and Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def text_size(draw: ImageDraw.ImageDraw, text: str, font) -> tuple[int, int]:
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def draw_text(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font,
    fill: str = "#1F2933",
    anchor: str | None = None,
):
    draw.text(xy, text, font=font, fill=fill, anchor=anchor)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = str(text).split()
    lines: list[str] = []
    current: list[str] = []
    for word in words:
        trial = " ".join(current + [word])
        if text_size(draw, trial, font)[0] <= max_width or not current:
            current.append(word)
        else:
            lines.append(" ".join(current))
            current = [word]
    if current:
        lines.append(" ".join(current))
    return lines


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    x: int,
    y: int,
    text: str,
    font,
    max_width: int,
    fill: str = "#1F2933",
    line_gap: int = 5,
):
    for line in wrap_text(draw, text, font, max_width):
        draw_text(draw, (x, y), line, font, fill)
        y += text_size(draw, line, font)[1] + line_gap
    return y


def save_image(img: Image.Image, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "PNG", optimize=True)


def draw_header(draw: ImageDraw.ImageDraw, title: str, subtitle: str, width: int):
    title_font = load_font(34, bold=True)
    subtitle_font = load_font(20)
    draw_text(draw, (80, 45), title, title_font, "#17324D")
    draw_wrapped(draw, 80, 92, subtitle, subtitle_font, width - 160, "#4D5B68", 6)


def scale_x(value: float, min_value: float, max_value: float, left: int, right: int) -> int:
    return int(left + ((value - min_value) / (max_value - min_value)) * (right - left))


def generate_construct_profile(construct_summary: pd.DataFrame) -> Path:
    path = FIGURES_DIR / "construct_mean_profile.png"
    width, height = 1800, 1180
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw_header(
        draw,
        "Construct Mean Profile",
        "Weighted means plotted against the four-point Likert scale with interpretation bands.",
        width,
    )

    label_font = load_font(20)
    value_font = load_font(19, bold=True)
    axis_font = load_font(18)
    left, right, top = 510, 1660, 170
    row_h, bar_h = 72, 28
    band_colors = [("#F7DAD6", 1.00, 1.75), ("#F7E6C4", 1.76, 2.50), ("#DDEEDB", 2.51, 3.25), ("#CFE4F3", 3.26, 4.00)]
    data = construct_summary.copy()

    for color, start, end in band_colors:
        draw.rectangle([scale_x(start, 1, 4, left, right), top - 18, scale_x(end, 1, 4, left, right), top + row_h * len(data) + 8], fill=color)
    for tick in [1, 2, 3, 4]:
        x = scale_x(tick, 1, 4, left, right)
        draw.line([x, top - 28, x, top + row_h * len(data) + 16], fill="#AEB8C2", width=2)
        draw_text(draw, (x, top + row_h * len(data) + 35), str(tick), axis_font, "#394B59", anchor="mm")

    for i, row in data.iterrows():
        y = top + i * row_h
        draw_wrapped(draw, 80, y - 10, row["Construct"], label_font, 400, "#1F2933", 4)
        mean = float(row["Mean"])
        x0 = scale_x(1, 1, 4, left, right)
        x1 = scale_x(mean, 1, 4, left, right)
        draw.rounded_rectangle([x0, y, x1, y + bar_h], radius=12, fill="#2F6F9F")
        draw.ellipse([x1 - 10, y - 4, x1 + 10, y + bar_h + 4], fill="#17324D")
        draw_text(draw, (x1 + 18, y + bar_h // 2), fmt(mean), value_font, "#17324D", anchor="lm")

    legend_y = height - 120
    legend_labels = ["Strongly Disagree / Very Low", "Disagree / Low", "Agree / High", "Strongly Agree / Very High"]
    x = 80
    for (color, _, _), label in zip(band_colors, legend_labels):
        draw.rectangle([x, legend_y, x + 28, legend_y + 22], fill=color, outline="#93A4B3")
        draw_text(draw, (x + 38, legend_y - 1), label, axis_font, "#394B59")
        x += text_size(draw, label, axis_font)[0] + 80

    save_image(img, path)
    return path


def generate_credibility_strategy_comparison(construct_summary: pd.DataFrame) -> Path:
    path = FIGURES_DIR / "credibility_strategy_comparison.png"
    width, height = 1800, 980
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw_header(
        draw,
        "Credibility and Communication Strategy Comparison",
        "Mean scores of primary credibility constructs compared with communication strategy outcomes.",
        width,
    )

    labels = ["Expertise", "Trustworthiness", "Overall Credibility", "Quality of Content", "Message Consistency", "Audience Engagement", "Overall Communication Strategy"]
    values = [float(construct_summary.loc[construct_summary["Construct"] == label, "Mean"].iloc[0]) for label in labels]
    colors = ["#2F6F9F", "#2F6F9F", "#17324D", "#5C946E", "#5C946E", "#5C946E", "#2F6F5E"]
    left, right, bottom, top = 160, 1680, 800, 190
    axis_font = load_font(18)
    label_font = load_font(18)
    value_font = load_font(19, bold=True)

    draw.line([left, bottom, right, bottom], fill="#52616D", width=3)
    draw.line([left, top, left, bottom], fill="#52616D", width=3)
    for tick in [1, 2, 3, 4]:
        y = int(bottom - ((tick - 1) / 3) * (bottom - top))
        draw.line([left - 8, y, right, y], fill="#D8DEE4", width=2)
        draw_text(draw, (left - 25, y), str(tick), axis_font, "#394B59", anchor="rm")

    slot = (right - left) / len(labels)
    bar_w = int(slot * 0.52)
    for i, (label, value, color) in enumerate(zip(labels, values, colors)):
        cx = int(left + slot * i + slot / 2)
        y = int(bottom - ((value - 1) / 3) * (bottom - top))
        draw.rounded_rectangle([cx - bar_w // 2, y, cx + bar_w // 2, bottom], radius=14, fill=color)
        draw_text(draw, (cx, y - 22), fmt(value), value_font, "#17324D", anchor="mm")
        for j, line in enumerate(wrap_text(draw, label, label_font, int(slot * 0.85))):
            draw_text(draw, (cx, bottom + 28 + j * 24), line, label_font, "#1F2933", anchor="mm")

    draw_text(draw, (100, 150), "Likert mean", axis_font, "#394B59")
    save_image(img, path)
    return path


def generate_reliability_chart(reliability: pd.DataFrame) -> Path:
    path = FIGURES_DIR / "reliability_alpha_chart.png"
    width, height = 1800, 1120
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw_header(
        draw,
        "Internal Consistency Reliability",
        "Cronbach's alpha coefficients with conventional reference thresholds for scale reliability.",
        width,
    )

    label_font = load_font(20)
    value_font = load_font(19, bold=True)
    axis_font = load_font(18)
    left, right, top = 570, 1660, 175
    row_h, bar_h = 72, 28
    data = reliability.copy()

    threshold_colors = [("#F7DAD6", 0.00, 0.59), ("#F7E6C4", 0.60, 0.69), ("#DDEEDB", 0.70, 0.79), ("#CFE4F3", 0.80, 0.89), ("#D8E8D4", 0.90, 1.00)]
    for color, start, end in threshold_colors:
        draw.rectangle([scale_x(start, 0, 1, left, right), top - 18, scale_x(end, 0, 1, left, right), top + row_h * len(data) + 10], fill=color)
    for tick in [0, 0.6, 0.7, 0.8, 0.9, 1.0]:
        x = scale_x(tick, 0, 1, left, right)
        draw.line([x, top - 28, x, top + row_h * len(data) + 16], fill="#AEB8C2", width=2)
        draw_text(draw, (x, top + row_h * len(data) + 35), f"{tick:.1f}", axis_font, "#394B59", anchor="mm")

    for i, row in data.iterrows():
        y = top + i * row_h
        draw_wrapped(draw, 80, y - 10, row["Scale"], label_font, 455, "#1F2933", 4)
        alpha = float(row["Cronbach Alpha"])
        color = "#17324D" if alpha >= 0.8 else "#2F6F9F" if alpha >= 0.7 else "#B7791F" if alpha >= 0.6 else "#A23B3B"
        x0 = scale_x(0, 0, 1, left, right)
        x1 = scale_x(alpha, 0, 1, left, right)
        draw.rounded_rectangle([x0, y, x1, y + bar_h], radius=12, fill=color)
        draw_text(draw, (x1 + 18, y + bar_h // 2), f"alpha = {fmt(alpha)}", value_font, "#17324D", anchor="lm")

    save_image(img, path)
    return path


def heat_color(value: float) -> tuple[int, int, int]:
    value = max(-1, min(1, value))
    if value >= 0:
        base = np.array([236, 246, 239])
        target = np.array([47, 111, 94])
        mix = value
    else:
        base = np.array([249, 235, 235])
        target = np.array([162, 59, 59])
        mix = abs(value)
    color = (base * (1 - mix) + target * mix).astype(int)
    return tuple(int(c) for c in color)


def generate_correlation_heatmap(correlations: pd.DataFrame) -> Path:
    path = FIGURES_DIR / "correlation_heatmap.png"
    width, height = 1700, 1040
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw_header(
        draw,
        "Correlation Heatmap",
        "Pearson r values between AI credibility predictors and communication strategy outcomes.",
        width,
    )

    predictors = ["Expertise", "Trustworthiness", "Overall Credibility"]
    outcomes = ["Quality of Content", "Message Consistency", "Audience Engagement", "Overall Communication Strategy"]
    left, top = 460, 250
    cell_w, cell_h = 280, 170
    label_font = load_font(20, bold=True)
    small_font = load_font(18)
    value_font = load_font(28, bold=True)

    for j, outcome in enumerate(outcomes):
        x = left + j * cell_w
        draw_wrapped(draw, x + 18, top - 85, outcome, small_font, cell_w - 35, "#1F2933", 4)
    for i, predictor in enumerate(predictors):
        y = top + i * cell_h
        draw_wrapped(draw, 80, y + 42, predictor, label_font, 330, "#17324D", 4)
        for j, outcome in enumerate(outcomes):
            x = left + j * cell_w
            row = correlations[(correlations["Predictor"] == predictor) & (correlations["Outcome"] == outcome)].iloc[0]
            r = float(row["r"])
            p = float(row["p-value"])
            draw.rounded_rectangle([x, y, x + cell_w - 12, y + cell_h - 12], radius=18, fill=heat_color(r), outline="#FFFFFF", width=4)
            sig = "*" if p < 0.05 else "ns"
            draw_text(draw, (x + cell_w // 2 - 6, y + 58), f"{r:.3f}", value_font, "#0E2538", anchor="mm")
            draw_text(draw, (x + cell_w // 2 - 6, y + 98), sig, small_font, "#394B59", anchor="mm")

    draw_text(draw, (80, height - 95), "* p < .05; ns = not significant. Darker green indicates stronger positive association.", small_font, "#394B59")
    save_image(img, path)
    return path


def generate_scatter(scored: pd.DataFrame, correlations: pd.DataFrame) -> Path:
    path = FIGURES_DIR / "credibility_strategy_scatter.png"
    width, height = 1700, 1080
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    focal = correlations[
        (correlations["Predictor"] == "Overall Credibility")
        & (correlations["Outcome"] == "Overall Communication Strategy")
    ].iloc[0]
    draw_header(
        draw,
        "Overall Credibility and Communication Strategy",
        "Respondent-level construct scores with fitted linear trend line.",
        width,
    )

    x_values = scored["Overall Credibility"].astype(float).to_numpy()
    y_values = scored["Overall Communication Strategy"].astype(float).to_numpy()
    left, right, top, bottom = 210, 1550, 185, 850
    axis_font = load_font(20)
    label_font = load_font(22, bold=True)
    value_font = load_font(24, bold=True)

    draw.rectangle([left, top, right, bottom], outline="#52616D", width=3)
    for tick in [2.0, 2.5, 3.0, 3.5, 4.0]:
        x = scale_x(tick, 2, 4, left, right)
        y = int(bottom - ((tick - 2) / 2) * (bottom - top))
        draw.line([x, top, x, bottom + 8], fill="#D8DEE4", width=2)
        draw.line([left - 8, y, right, y], fill="#D8DEE4", width=2)
        draw_text(draw, (x, bottom + 35), f"{tick:.1f}", axis_font, "#394B59", anchor="mm")
        draw_text(draw, (left - 28, y), f"{tick:.1f}", axis_font, "#394B59", anchor="rm")

    for x_val, y_val in zip(x_values, y_values):
        x = scale_x(float(x_val), 2, 4, left, right)
        y = int(bottom - ((float(y_val) - 2) / 2) * (bottom - top))
        draw.ellipse([x - 8, y - 8, x + 8, y + 8], fill="#2F6F9F", outline="#17324D")

    slope, intercept = np.polyfit(x_values, y_values, 1)
    x1, x2 = x_values.min(), x_values.max()
    y1, y2 = slope * x1 + intercept, slope * x2 + intercept
    draw.line(
        [
            scale_x(float(x1), 2, 4, left, right),
            int(bottom - ((float(y1) - 2) / 2) * (bottom - top)),
            scale_x(float(x2), 2, 4, left, right),
            int(bottom - ((float(y2) - 2) / 2) * (bottom - top)),
        ],
        fill="#A23B3B",
        width=6,
    )
    draw_text(draw, ((left + right) // 2, bottom + 85), "Overall AI Credibility", label_font, "#17324D", anchor="mm")
    draw_text(draw, (left, top - 38), "Overall Communication Strategy", label_font, "#17324D", anchor="lm")
    p_text = "< .001" if float(focal["p-value"]) < 0.001 else fmt(float(focal["p-value"]))
    annotation = f"r = {float(focal['r']):.3f} | p = {p_text} | n = {int(focal['n'])}"
    draw.rounded_rectangle([1010, 220, 1505, 310], radius=18, fill="#F0F5FA", outline="#B8C7D3", width=2)
    draw_text(draw, (1035, 250), annotation, value_font, "#17324D")

    save_image(img, path)
    return path


def generate_figures(
    scored: pd.DataFrame,
    construct_summary: pd.DataFrame,
    reliability: pd.DataFrame,
    correlations: pd.DataFrame,
) -> dict[str, Path]:
    FIGURES_DIR.mkdir(exist_ok=True)
    return {
        "construct_profile": generate_construct_profile(construct_summary),
        "credibility_strategy": generate_credibility_strategy_comparison(construct_summary),
        "reliability": generate_reliability_chart(reliability),
        "correlation_heatmap": generate_correlation_heatmap(correlations),
        "scatter": generate_scatter(scored, correlations),
    }


def add_run(paragraph, text: str, bold: bool = False, italic: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    return run


def add_df_table(document: Document, df: pd.DataFrame, columns: list[str], title: str | None = None):
    if title:
        p = document.add_paragraph()
        add_run(p, title, bold=True)
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, col in enumerate(columns):
        hdr[i].text = col
        for paragraph in hdr[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(columns):
            val = row[col]
            if isinstance(val, float):
                if col.lower().startswith("p"):
                    text = "< .001" if val < 0.001 else fmt(val, 3)
                else:
                    text = fmt(val, 3)
            else:
                text = str(val)
            cells[i].text = text
    document.add_paragraph()
    return table


def add_report_figure(document: Document, path: Path, caption: str, interpretation: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(6.35))

    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(cap, caption, bold=True)
    for run in cap.runs:
        run.font.size = Pt(9)

    interp = document.add_paragraph(interpretation)
    interp.paragraph_format.space_after = Pt(8)


def set_doc_styles(document: Document):
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.size = Pt(12)


def add_footer(document: Document):
    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "AITHOS Statistical Analysis Report | Computed by Engr. Jamie Eduardo Rosal, MSCpE"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.name = "Arial"


def write_report(
    scored: pd.DataFrame,
    item_summary: pd.DataFrame,
    construct_summary: pd.DataFrame,
    reliability: pd.DataFrame,
    correlations: pd.DataFrame,
    validation_messages: list[str],
    figure_paths: dict[str, Path],
):
    doc = Document()
    set_doc_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Statistical Analysis Report")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(20)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(
        subtitle,
        "The Effect of Credibility of AI-Generated Business Content on Communication Strategies "
        "Among Coffee Shop Owners and Managers in CAMANAVA",
        bold=True,
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("\nComputed by\n").bold = True
    meta.add_run("Engr. Jamie Eduardo Rosal, MSCpE\n").bold = True
    meta.add_run("May 2026")

    doc.add_page_break()

    doc.add_heading("Abstract / Executive Statistical Summary", level=1)
    n = len(scored)
    overall_cred = construct_summary.loc[construct_summary["Construct"] == "Overall Credibility"].iloc[0]
    overall_comm = construct_summary.loc[construct_summary["Construct"] == "Overall Communication Strategy"].iloc[0]
    focal = correlations[
        (correlations["Predictor"] == "Overall Credibility")
        & (correlations["Outcome"] == "Overall Communication Strategy")
    ].iloc[0]
    doc.add_paragraph(
        f"This report presents a quantitative statistical analysis of {n} valid survey responses from the "
        "available AITHOS survey workbook. The analysis is aligned with the research problem stated in the "
        "final research manuscript and the structured instrument on AI-generated business content credibility, "
        "algorithmic transparency, technology acceptance, and communication strategies. Because the provided "
        "workbook contains Parts II to V only, Part I business profile statistics are not computed and no "
        "profile data are inferred."
    )
    doc.add_paragraph(
        f"The respondents generally evaluated AI-generated business content favorably. Overall AI credibility "
        f"obtained a mean of {fmt(overall_cred['Mean'])} ({overall_cred['Interpretation']}), while overall "
        f"communication strategy obtained a mean of {fmt(overall_comm['Mean'])} "
        f"({overall_comm['Interpretation']}). The primary inferential result showed a {focal['Interpretation']} "
        f"relationship between overall credibility and overall communication strategy, r = {fmt(focal['r'])}, "
        f"p = {'< .001' if focal['p-value'] < 0.001 else fmt(focal['p-value'])}, indicating that the relationship "
        f"is {focal['Decision'].lower()} at alpha = {ALPHA}."
    )

    doc.add_heading("Methodology", level=1)
    doc.add_heading("Dataset and Scope", level=2)
    doc.add_paragraph(
        "The official source dataset is AITHOS-SURVEY-RAW-DATA.xlsx. The workbook contains four usable sheets: "
        "Part II for credibility, Part III for algorithmic transparency, Part IV for technology acceptance, and "
        "Part V for communication strategies. The dataset contains 60 aligned respondent records across all "
        "available construct sheets."
    )
    doc.add_paragraph(
        "The final research manuscript also calls for a respondent business profile analysis, but the available "
        "workbook does not include the Part I profile variables. Accordingly, this report omits the business "
        "profile tables and limits the computations to the available Likert-scale data. This limitation protects "
        "the integrity of the analysis by avoiding invented or unverified profile values."
    )

    doc.add_heading("Instrument and Scoring Procedure", level=2)
    doc.add_paragraph(
        "The survey instrument uses a four-point Likert scale: 4 = Strongly Agree, 3 = Agree, 2 = Disagree, "
        "and 1 = Strongly Disagree. Item means were interpreted as follows: 3.26-4.00 = Strongly Agree / Very "
        "High, 2.51-3.25 = Agree / High, 1.76-2.50 = Disagree / Low, and 1.00-1.75 = Strongly Disagree / Very Low."
    )
    doc.add_paragraph(
        "For each construct, respondent-level scores were computed as the arithmetic mean of the five item "
        "responses under that construct. Overall AI Credibility was computed from Expertise and Trustworthiness. "
        "Overall Communication Strategy was computed from Quality of Content, Message Consistency, and Audience "
        "Engagement. Technology Acceptance scores were summarized descriptively as supplemental context."
    )

    doc.add_heading("Statistical Treatments", level=2)
    doc.add_paragraph(
        "The analysis used descriptive statistics, Cronbach's alpha reliability coefficients, and Pearson product-"
        "moment correlations. Pearson correlation was selected as the primary inferential treatment because the "
        "research question asks whether AI credibility is significantly related to communication strategies. "
        "Statistical significance was evaluated at alpha = 0.05."
    )

    doc.add_heading("Data Validation Summary", level=2)
    for message in validation_messages:
        doc.add_paragraph(message, style=None)

    doc.add_heading("Results", level=1)
    doc.add_heading("Descriptive Statistics by Construct", level=2)
    add_df_table(
        doc,
        construct_summary,
        ["Construct", "Mean", "SD", "Min", "Max", "Interpretation"],
        "Table 1. Construct-Level Descriptive Statistics",
    )

    doc.add_paragraph(
        "The construct-level descriptive statistics show that the available respondent ratings cluster in the "
        "Agree to Strongly Agree range. This suggests that AI-generated business content is generally viewed as "
        "credible and useful by the sampled coffee shop owners and managers, while communication strategy ratings "
        "also indicate favorable perceived performance."
    )

    doc.add_heading("Item-Level Weighted Means", level=2)
    for construct in CONSTRUCTS:
        subset = item_summary[item_summary["Construct"] == construct][["Item", "Statement", "Mean", "SD", "Interpretation"]]
        add_df_table(doc, subset, ["Item", "Statement", "Mean", "SD", "Interpretation"], f"Table. {construct} Items")

    doc.add_heading("Visual Analysis of Findings", level=2)
    add_report_figure(
        doc,
        figure_paths["construct_profile"],
        "Figure 1. Construct mean profile plotted against the four-point Likert scale.",
        "Figure 1 shows that the construct means generally fall within the Agree / High and Strongly Agree / "
        "Very High interpretation ranges. This visual profile supports the descriptive conclusion that the "
        "respondents evaluated AI credibility, technology acceptance, and AI-supported communication strategies "
        "favorably, while still preserving differences between constructs.",
    )
    add_report_figure(
        doc,
        figure_paths["credibility_strategy"],
        "Figure 2. Mean comparison of AI credibility and communication strategy constructs.",
        "Figure 2 places the primary independent and dependent constructs side by side. The communication strategy "
        "subscales, particularly message consistency and audience engagement, are positioned near or above the "
        "Strongly Agree threshold, while credibility indicators remain within the Agree / High range.",
    )
    add_report_figure(
        doc,
        figure_paths["reliability"],
        "Figure 3. Cronbach's alpha reliability coefficients by scale.",
        "Figure 3 highlights the internal consistency of the instrument scales. Overall AI Credibility shows good "
        "reliability, several individual scales are acceptable, and lower alpha values are made visible so readers "
        "can interpret specific subscales with appropriate caution.",
    )
    add_report_figure(
        doc,
        figure_paths["correlation_heatmap"],
        "Figure 4. Heatmap of Pearson correlations between credibility predictors and strategy outcomes.",
        "Figure 4 makes the inferential pattern easier to inspect: credibility is most strongly associated with "
        "quality of content, while audience engagement shows weaker and non-significant relationships. This "
        "distinction is useful because it shows that credibility does not affect every communication outcome with "
        "the same magnitude.",
    )
    add_report_figure(
        doc,
        figure_paths["scatter"],
        "Figure 5. Respondent-level relationship between overall credibility and overall communication strategy.",
        "Figure 5 provides a respondent-level view of the primary test. The upward trend line and the annotated "
        "coefficient show a significant moderate positive relationship, indicating that higher perceived AI "
        "credibility tends to accompany stronger overall communication strategy ratings.",
    )

    doc.add_heading("Reliability Analysis", level=2)
    add_df_table(
        doc,
        reliability,
        ["Scale", "Items", "Cronbach Alpha", "Interpretation"],
        "Table 2. Internal Consistency Reliability",
    )
    doc.add_paragraph(
        "Cronbach's alpha was used to evaluate whether the items within each multi-item scale consistently "
        "measure the same latent construct. Higher alpha coefficients indicate stronger internal consistency, "
        "while lower values should be interpreted as evidence that items may be measuring a more heterogeneous "
        "set of perceptions."
    )

    doc.add_heading("Correlation Analysis", level=2)
    add_df_table(
        doc,
        correlations,
        ["Predictor", "Outcome", "n", "r", "p-value", "Decision", "Interpretation"],
        "Table 3. Pearson Correlations Between AI Credibility and Communication Strategies",
    )
    doc.add_paragraph(
        "The primary test addresses whether AI credibility is significantly associated with communication "
        "strategies among CAMANAVA coffee shop owners and managers. The correlation table should be read by "
        "examining the sign, magnitude, and significance of r. Positive coefficients indicate that higher "
        "credibility ratings tend to move with stronger communication strategy ratings."
    )

    doc.add_heading("Discussion", level=1)
    doc.add_paragraph(
        "The findings support the conceptual logic of Source Credibility Theory in the AI-generated business "
        "content context. When AI content is perceived as knowledgeable, professional, dependable, and aligned "
        "with brand communication needs, coffee shop owners and managers also tend to evaluate their AI-assisted "
        "communication strategies more favorably. This is theoretically consistent with the proposition that "
        "credibility strengthens persuasion and acceptance of communication sources."
    )
    doc.add_paragraph(
        "Algorithmic transparency and TAM3 variables were treated as descriptive supplemental constructs in this "
        "report. Their inclusion remains important because transparency, usefulness, and ease of use provide "
        "context for why business users may accept or reject AI-generated content in daily marketing communication. "
        "However, the primary inferential emphasis remains the relationship between AI credibility and communication "
        "strategies, as specified in the approved analysis plan."
    )
    doc.add_paragraph(
        "For CAMANAVA coffee shops, the results imply that AI adoption should not be evaluated merely by speed or "
        "automation convenience. Owners and managers should also examine whether AI-generated messages remain "
        "accurate, trustworthy, brand-consistent, locally understandable, and capable of supporting customer "
        "engagement without weakening perceived authenticity."
    )

    doc.add_heading("Conclusion and Recommendations", level=1)
    doc.add_paragraph(
        "Based on the available survey data, AI-generated business content credibility is positively associated "
        "with communication strategy outcomes. This indicates that strengthening the perceived expertise and "
        "trustworthiness of AI-generated content may contribute to better content quality, stronger message "
        "consistency, and improved audience engagement."
    )
    recommendations = [
        "Coffee shop owners and managers should review AI-generated content before publication to preserve accuracy, warmth, and brand voice.",
        "AI tools should be used as communication support systems rather than full replacements for human judgment.",
        "Businesses should document prompt practices, review standards, and content approval workflows to improve credibility and transparency.",
        "Future research should collect and analyze Part I profile variables to test whether business size, AI usage duration, and frequency of use shape credibility perceptions.",
    ]
    for rec in recommendations:
        doc.add_paragraph(rec, style="List Bullet")

    doc.add_heading("Appendix A: Construct Mapping", level=1)
    mapping_rows = []
    for construct, spec in CONSTRUCTS.items():
        mapping_rows.append(
            {
                "Construct": construct,
                "Domain": spec["domain"],
                "Source Sheet": spec["sheet"],
                "Items": len(spec["items"]),
            }
        )
    add_df_table(doc, pd.DataFrame(mapping_rows), ["Construct", "Domain", "Source Sheet", "Items"])

    doc.add_heading("Appendix B: Validation Notebook", level=1)
    doc.add_paragraph(
        f"The reproducible computation file is {NOTEBOOK_FILE.name}. It loads the raw workbook, reconstructs the "
        "construct mapping, validates the Likert response range and respondent alignment, and recomputes all "
        "descriptive, reliability, and correlation tables used in this report."
    )

    add_footer(doc)
    doc.save(REPORT_FILE)


def make_notebook_source() -> list[dict]:
    code = r'''
from pathlib import Path
import math
import numpy as np
import pandas as pd

ROOT = Path(".").resolve()
DATA_FILE = ROOT / "AITHOS-SURVEY-RAW-DATA.xlsx"
ALPHA = 0.05

CONSTRUCTS = {
    "Expertise": {"sheet": "PART II.", "cols": [1, 2, 3, 4, 5]},
    "Trustworthiness": {"sheet": "PART II.", "cols": [7, 8, 9, 10, 11]},
    "Algorithmic Transparency": {"sheet": "PART III.", "cols": [1, 2, 3, 4, 5]},
    "Perceived Usefulness": {"sheet": "PART IV.", "cols": [1, 2, 3, 4, 5]},
    "Perceived Ease of Use": {"sheet": "PART IV.", "cols": [7, 8, 9, 10, 11]},
    "Quality of Content": {"sheet": "PART V.", "cols": [1, 2, 3, 4, 5]},
    "Message Consistency": {"sheet": "PART V.", "cols": [7, 8, 9, 10, 11]},
    "Audience Engagement": {"sheet": "PART V.", "cols": [14, 15, 16, 17, 18]},
}

def verbal_interpretation(mean):
    if mean >= 3.26:
        return "Strongly Agree / Very High"
    if mean >= 2.51:
        return "Agree / High"
    if mean >= 1.76:
        return "Disagree / Low"
    return "Strongly Disagree / Very Low"

def alpha_interpretation(alpha):
    if pd.isna(alpha):
        return "Not computable"
    if alpha >= 0.90:
        return "Excellent"
    if alpha >= 0.80:
        return "Good"
    if alpha >= 0.70:
        return "Acceptable"
    if alpha >= 0.60:
        return "Questionable"
    return "Low"

def correlation_interpretation(r):
    ar = abs(r)
    if ar >= 0.90:
        strength = "Very strong"
    elif ar >= 0.70:
        strength = "Strong"
    elif ar >= 0.50:
        strength = "Moderate"
    elif ar >= 0.30:
        strength = "Weak"
    elif ar >= 0.10:
        strength = "Very weak"
    else:
        strength = "Negligible"
    return f"{strength} {'positive' if r >= 0 else 'negative'}"

def p_value_from_r(r, n):
    if n <= 3 or not np.isfinite(r):
        return np.nan
    r = min(max(float(r), -0.999999999), 0.999999999)
    z = math.atanh(r) * math.sqrt(n - 3)
    return math.erfc(abs(z) / math.sqrt(2.0))

def cronbach_alpha(df):
    clean = df.dropna()
    k = clean.shape[1]
    if k < 2 or clean.shape[0] < 2:
        return np.nan
    item_variances = clean.var(axis=0, ddof=1)
    total_variance = clean.sum(axis=1).var(ddof=1)
    if total_variance == 0:
        return np.nan
    return (k / (k - 1)) * (1 - item_variances.sum() / total_variance)

raw = pd.read_excel(DATA_FILE, sheet_name=None, header=None)
construct_items = {}
respondent_ids = None

for name, spec in CONSTRUCTS.items():
    sheet = raw[spec["sheet"]]
    block = sheet.iloc[9:, [0] + spec["cols"]].copy()
    block.columns = ["Respondent"] + [f"{name} {i}" for i in range(1, 6)]
    block["Respondent"] = pd.to_numeric(block["Respondent"], errors="coerce").astype("Int64")
    for col in block.columns[1:]:
        block[col] = pd.to_numeric(block[col], errors="coerce")
    block = block.dropna(subset=["Respondent"]).set_index("Respondent").sort_index()
    construct_items[name] = block
    if respondent_ids is None:
        respondent_ids = block.index
    else:
        assert respondent_ids.equals(block.index), f"Respondent IDs do not align for {name}"

scored = pd.DataFrame(index=respondent_ids)
for name, block in construct_items.items():
    scored[name] = block.mean(axis=1)

scored["Overall Credibility"] = scored[["Expertise", "Trustworthiness"]].mean(axis=1)
scored["Overall Technology Acceptance"] = scored[["Perceived Usefulness", "Perceived Ease of Use"]].mean(axis=1)
scored["Overall Communication Strategy"] = scored[["Quality of Content", "Message Consistency", "Audience Engagement"]].mean(axis=1)

all_values = pd.concat(construct_items.values(), axis=1)
print("Respondents:", len(scored))
print("Respondent ID range:", int(scored.index.min()), "to", int(scored.index.max()))
print("Out-of-range Likert values:", int(((all_values < 1) | (all_values > 4)).sum().sum()))
print("Missing Likert values:", int(all_values.isna().sum().sum()))
for name, block in construct_items.items():
    print(f"{name}: {block.shape[1]} items x {block.shape[0]} respondents")
'''
    desc_code = r'''
construct_order = [
    "Expertise", "Trustworthiness", "Overall Credibility",
    "Algorithmic Transparency",
    "Perceived Usefulness", "Perceived Ease of Use", "Overall Technology Acceptance",
    "Quality of Content", "Message Consistency", "Audience Engagement", "Overall Communication Strategy",
]

construct_summary = pd.DataFrame([
    {
        "Construct": name,
        "Mean": scored[name].mean(),
        "SD": scored[name].std(ddof=1),
        "Min": scored[name].min(),
        "Max": scored[name].max(),
        "Interpretation": verbal_interpretation(scored[name].mean()),
    }
    for name in construct_order
])
construct_summary
'''
    item_code = r'''
item_rows = []
for construct, block in construct_items.items():
    for idx, col in enumerate(block.columns, start=1):
        item_rows.append({
            "Construct": construct,
            "Item": idx,
            "Mean": block[col].mean(),
            "SD": block[col].std(ddof=1),
            "Interpretation": verbal_interpretation(block[col].mean()),
        })
item_summary = pd.DataFrame(item_rows)
item_summary
'''
    reliability_code = r'''
reliability_rows = []
for construct, block in construct_items.items():
    alpha = cronbach_alpha(block)
    reliability_rows.append({
        "Scale": construct,
        "Items": block.shape[1],
        "Cronbach Alpha": alpha,
        "Interpretation": alpha_interpretation(alpha),
    })

combined = {
    "Overall AI Credibility": ["Expertise", "Trustworthiness"],
    "Overall Communication Strategy": ["Quality of Content", "Message Consistency", "Audience Engagement"],
    "Overall Technology Acceptance": ["Perceived Usefulness", "Perceived Ease of Use"],
}
for label, names in combined.items():
    block = pd.concat([construct_items[name] for name in names], axis=1)
    alpha = cronbach_alpha(block)
    reliability_rows.append({
        "Scale": label,
        "Items": block.shape[1],
        "Cronbach Alpha": alpha,
        "Interpretation": alpha_interpretation(alpha),
    })

reliability = pd.DataFrame(reliability_rows)
reliability
'''
    corr_code = r'''
correlation_rows = []
predictors = ["Expertise", "Trustworthiness", "Overall Credibility"]
outcomes = ["Quality of Content", "Message Consistency", "Audience Engagement", "Overall Communication Strategy"]

for predictor in predictors:
    for outcome in outcomes:
        pair = scored[[predictor, outcome]].dropna()
        n = len(pair)
        r = pair[predictor].corr(pair[outcome])
        p = p_value_from_r(r, n)
        correlation_rows.append({
            "Predictor": predictor,
            "Outcome": outcome,
            "n": n,
            "r": r,
            "p-value": p,
            "Decision": "Significant" if p < ALPHA else "Not significant",
            "Interpretation": correlation_interpretation(r),
        })

correlations = pd.DataFrame(correlation_rows)
correlations
'''
    viz_code = r'''
from build_aithos_analysis import generate_figures

figure_paths = generate_figures(scored, construct_summary, reliability, correlations)
for name, path in figure_paths.items():
    print(f"{name}: {path}")

try:
    from IPython.display import Image, display
    for path in figure_paths.values():
        display(Image(filename=str(path)))
except Exception:
    print("Figures generated. Inline display is available when this notebook is opened in Jupyter.")
'''
    export_code = r'''
with pd.ExcelWriter("AITHOS_Computation_Tables.xlsx", engine="openpyxl") as writer:
    scored.reset_index().to_excel(writer, sheet_name="Respondent Scores", index=False)
    construct_summary.to_excel(writer, sheet_name="Construct Summary", index=False)
    item_summary.to_excel(writer, sheet_name="Item Summary", index=False)
    reliability.to_excel(writer, sheet_name="Reliability", index=False)
    correlations.to_excel(writer, sheet_name="Correlations", index=False)

print("Exported AITHOS_Computation_Tables.xlsx for audit/reference.")
'''
    return [
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": [
                "# AITHOS Statistical Computations Validation\n\n",
                "This notebook validates the computations for the AITHOS statistical analysis report. ",
                "It uses the raw workbook as the authoritative data source and does not modify the raw file.\n\n",
                "**Computed by:** Engr. Jamie Eduardo Rosal, MSCpE\n",
            ],
        },
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": [
                "## 1. Load Data, Define Constructs, and Validate Responses\n\n",
                "The workbook contains Parts II-V only. Part I profile variables are unavailable and are therefore omitted from computation.\n",
            ],
        },
        {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [], "source": code.strip().splitlines(True)},
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": ["## 2. Descriptive Statistics by Construct\n\nWeighted means are computed as arithmetic means of Likert responses.\n"],
        },
        {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [], "source": desc_code.strip().splitlines(True)},
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": ["## 3. Item-Level Statistics\n\nEach item is summarized by mean, standard deviation, and verbal interpretation.\n"],
        },
        {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [], "source": item_code.strip().splitlines(True)},
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": ["## 4. Reliability Analysis\n\nCronbach's alpha estimates internal consistency for each multi-item construct.\n"],
        },
        {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [], "source": reliability_code.strip().splitlines(True)},
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": ["## 5. Pearson Correlation Analysis\n\nPrimary inferential analysis tests the relationship between AI credibility and communication strategy constructs.\n"],
        },
        {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [], "source": corr_code.strip().splitlines(True)},
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": [
                "## 6. Publication-Style Visualizations\n\n",
                "The following cell regenerates the same PNG figures inserted in the Word report.\n",
            ],
        },
        {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [], "source": viz_code.strip().splitlines(True)},
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": [
                "![Construct mean profile](figures/construct_mean_profile.png)\n\n",
                "![Credibility and communication strategy comparison](figures/credibility_strategy_comparison.png)\n\n",
                "![Cronbach alpha reliability chart](figures/reliability_alpha_chart.png)\n\n",
                "![Correlation heatmap](figures/correlation_heatmap.png)\n\n",
                "![Overall credibility and communication strategy scatter plot](figures/credibility_strategy_scatter.png)\n",
            ],
        },
        {
            "cell_type": "markdown",
            "metadata": {},
            "source": ["## 7. Optional Export of Computation Tables\n\nThis cell creates an Excel audit workbook from the computed tables.\n"],
        },
        {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [], "source": export_code.strip().splitlines(True)},
    ]


def write_notebook():
    notebook = {
        "cells": make_notebook_source(),
        "metadata": {
            "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
            "language_info": {
                "name": "python",
                "version": "3.11",
                "mimetype": "text/x-python",
                "codemirror_mode": {"name": "ipython", "version": 3},
                "pygments_lexer": "ipython3",
                "nbconvert_exporter": "python",
                "file_extension": ".py",
            },
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    NOTEBOOK_FILE.write_text(json.dumps(notebook, indent=2), encoding="utf-8")


def main():
    scored, construct_items = clean_constructs()
    validation_messages = validate_data(scored, construct_items)
    item_summary = summarize_items(construct_items)
    construct_summary = summarize_constructs(scored)
    reliability = summarize_reliability(construct_items)
    correlations = summarize_correlations(scored)
    figure_paths = generate_figures(scored, construct_summary, reliability, correlations)

    write_report(scored, item_summary, construct_summary, reliability, correlations, validation_messages, figure_paths)
    write_notebook()

    print(f"Wrote {REPORT_FILE}")
    print(f"Wrote {NOTEBOOK_FILE}")
    for path in figure_paths.values():
        print(f"Wrote {path}")
    print("\nKey result:")
    focal = correlations[
        (correlations["Predictor"] == "Overall Credibility")
        & (correlations["Outcome"] == "Overall Communication Strategy")
    ].iloc[0]
    print(
        f"Overall Credibility vs Overall Communication Strategy: r={focal['r']:.3f}, "
        f"p={focal['p-value']:.3f}, {focal['Decision']}"
    )


if __name__ == "__main__":
    main()
