from __future__ import annotations

import math
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "outputs"
FIG_DIR = ROOT / "figures"
REPORT_PATH = ROOT / "Garlic_Peel_Powder_Full_Report_Rosal.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
HEADER_FILL = "F2F4F7"
LIGHT_FILL = "F7F9FC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D6DAE1")


def set_table_width(table, width_dxa=9360, indent_dxa=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def fmt(value, digits=3) -> str:
    if pd.isna(value):
        return "N/A"
    if isinstance(value, str):
        return value
    try:
        value = float(value)
    except Exception:
        return str(value)
    if math.isclose(value, round(value), abs_tol=1e-9):
        return str(int(round(value)))
    return f"{value:.{digits}f}"


def p_fmt(value) -> str:
    if pd.isna(value):
        return "N/A"
    value = float(value)
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def add_run(paragraph, text: str, bold=False, italic=False, color=None, size=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return run


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, text, italic=True, color=RGBColor(90, 98, 110), size=9)


def style_table(table, header=True, font_size=8.5) -> None:
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(font_size)
            if header and r_idx == 0:
                set_cell_shading(cell, HEADER_FILL)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = INK


def add_df_table(
    doc: Document,
    df: pd.DataFrame,
    title: str,
    note: str | None = None,
    digits: int = 3,
    font_size: float = 8.5,
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    add_run(p, title, bold=True, color=INK)
    if note:
        n = doc.add_paragraph()
        n.paragraph_format.space_after = Pt(4)
        add_run(n, note, italic=True, color=RGBColor(90, 98, 110), size=9)

    table = doc.add_table(rows=1, cols=len(df.columns))
    hdr = table.rows[0].cells
    for i, col in enumerate(df.columns):
        hdr[i].text = str(col)
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            val = row[col]
            col_l = str(col).lower()
            if col_l in {"p", "p-value", "p_value"} or col_l.startswith("p "):
                cells[i].text = p_fmt(val)
            else:
                cells[i].text = fmt(val, digits=digits)
    style_table(table, font_size=font_size)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "Full Analysis Report | Ilocos White Garlic Peel Powder | Reported by Engr. Jamie Eduardo Rosal, MSCpE", size=8, color=RGBColor(100, 108, 120))
    return doc


def add_title_page(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "FULL ANALYSIS REPORT", bold=True, color=BLUE, size=18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(
        p,
        "Antimicrobial Potential of Ilocos White Garlic (Allium sativum L.) Peel Powder Processed Under Varying Drying Time and Temperature Conditions",
        bold=True,
        color=INK,
        size=16,
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "Aligned with Research Objectives, Methodology, Statistical Analysis, and RSM Optimization", italic=True, color=RGBColor(90, 98, 110), size=11)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "Reported by", color=RGBColor(90, 98, 110), size=11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "Engr. Jamie Eduardo Rosal, MSCpE", bold=True, color=INK, size=14)
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "Source files: reference.docx and Book-of-Analysis-2026-rhodea.xlsx", color=RGBColor(90, 98, 110), size=9)
    doc.add_page_break()


def main() -> None:
    phys = pd.read_csv(OUT_DIR / "clean_physicochemical_data.csv")
    anti = pd.read_csv(OUT_DIR / "clean_antimicrobial_data.csv")
    anti_anova = pd.read_csv(OUT_DIR / "one_way_anova_antimicrobial.csv")
    anova_phys = pd.read_csv(OUT_DIR / "two_way_anova_physicochemical.csv")
    opt = pd.read_csv(OUT_DIR / "rsm_optimization_top_candidates.csv")

    treatment_map = pd.DataFrame(
        {
            "Treatment": range(1, 10),
            "Temperature (°C)": [50, 50, 50, 60, 60, 60, 70, 70, 70],
            "Time (h)": [4, 6, 8, 4, 6, 8, 4, 6, 8],
        }
    )

    doc = setup_document()
    add_title_page(doc)

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "This report presents the Python-based statistical analysis and visualization of Ilocos white garlic peel powder processed under varying oven-drying temperatures and times. The analysis was designed to directly address the study objectives: develop and compare treatment conditions, determine physicochemical characteristics, evaluate antimicrobial activity, and identify optimized processing conditions through Response Surface Methodology."
    )
    add_bullets(
        doc,
        [
            "The objectives were computationally achieved using the available workbook data.",
            "Treatment 5, corresponding to 60°C for 6 hours, was the best observed condition by composite desirability.",
            "The RSM screening suggested an optimized processing region near 57.33°C and 5.87 hours.",
            "Temperature, time, and their interaction significantly affected several physicochemical responses.",
            "For antimicrobial activity, treatment differences were significant for Staphylococcus aureus; Escherichia coli was not significant, while Salmonella spp. values were constant and therefore not suitable for meaningful ANOVA interpretation.",
        ],
    )

    doc.add_heading("Objective Alignment", level=1)
    objective_df = pd.DataFrame(
        [
            ["Develop garlic peel powder using oven-drying at varying time and temperature conditions.", "Achieved computationally", "Treatments were mapped into a 3x3 factorial design using temperature and time."],
            ["Determine physicochemical characteristics: Aw, moisture content, color, WAC, and WSI.", "Achieved", "Cleaned data were summarized by treatment and analyzed using two-way ANOVA."],
            ["Evaluate antimicrobial capacity against Staphylococcus aureus, Escherichia coli, and Salmonella spp.", "Achieved with data limitations", "Zone of inhibition data were summarized and analyzed using one-way ANOVA by organism."],
            ["Use RSM to generate response surfaces and identify optimized garlic peel powder conditions.", "Achieved", "Second-order response surface models and composite desirability screening were computed."],
        ],
        columns=["Research Objective", "Status", "Evidence in Report"],
    )
    add_df_table(doc, objective_df, "Table 1. Alignment of Computations with Study Objectives", font_size=8)

    doc.add_heading("Methodology and Data Preparation", level=1)
    doc.add_paragraph(
        "The study used a quantitative experimental design. The independent variables were oven-drying temperature and time; the dependent variables were physicochemical properties and antimicrobial zone of inhibition. The workbook was cleaned into analysis-ready datasets before performing descriptive statistics, ANOVA, and RSM."
    )
    doc.add_paragraph(
        "Note: the methodology text states drying times of 4, 6, and 8 hours, but the treatment matrix lists Treatment 9 as 70°C and 7 hours. For the statistical analysis, Treatment 9 was treated as 70°C and 8 hours to complete the stated 3x3 full factorial design. This should be corrected or clarified in the manuscript."
    )
    add_df_table(doc, treatment_map, "Table 2. Experimental Treatment Matrix Used for Analysis", digits=0)

    doc.add_heading("Physicochemical Results", level=1)
    phys_means = (
        phys.groupby(["Treatment", "Temperature_C", "Time_h"], as_index=False)
        [["WAC", "Aw", "MC_percent", "WSI", "L_star", "a_star", "b_star", "Chroma_C", "Delta_E_vs_T1"]]
        .mean()
    )
    phys_means = phys_means.rename(
        columns={
            "Temperature_C": "Temp (°C)",
            "Time_h": "Time (h)",
            "MC_percent": "MC (%)",
            "L_star": "L*",
            "a_star": "a*",
            "b_star": "b*",
            "Chroma_C": "C*",
            "Delta_E_vs_T1": "Delta E vs T1",
        }
    )
    add_df_table(
        doc,
        phys_means[["Treatment", "Temp (°C)", "Time (h)", "WAC", "Aw", "MC (%)", "WSI", "L*", "a*", "b*", "C*", "Delta E vs T1"]],
        "Table 3. Treatment Mean Physicochemical Characteristics",
        note="Values are means of triplicate measurements after workbook cleaning. Color C* and Delta E were recomputed because the workbook's derived color columns were zero-filled.",
        digits=3,
        font_size=7.2,
    )

    doc.add_picture(str(FIG_DIR / "physicochemical_heatmaps.png"), width=Inches(6.25))
    add_caption(doc, "Figure 1. Heatmaps of physicochemical treatment means across drying temperature and time.")
    doc.add_picture(str(FIG_DIR / "physicochemical_interaction_plots.png"), width=Inches(6.25))
    add_caption(doc, "Figure 2. Interaction plots for selected physicochemical responses.")

    doc.add_heading("Two-Way ANOVA Results", level=2)
    pivot_rows = []
    for response, sub in anova_phys.groupby("Response"):
        row = {"Response": response}
        for source, label in [
            ("C(Temperature_C)", "p Temperature"),
            ("C(Time_h)", "p Time"),
            ("C(Temperature_C):C(Time_h)", "p Interaction"),
        ]:
            vals = sub.loc[sub["Source"] == source, "PR(>F)"]
            row[label] = vals.iloc[0] if len(vals) else float("nan")
        pivot_rows.append(row)
    anova_pivot = pd.DataFrame(pivot_rows)
    add_df_table(
        doc,
        anova_pivot,
        "Table 4. Two-Way ANOVA p-Values for Physicochemical Responses",
        note="p < 0.05 indicates a statistically significant effect.",
        font_size=8,
    )

    doc.add_heading("Antimicrobial Results", level=1)
    anti_means = (
        anti[anti["Treatment"].between(1, 9)]
        .groupby(["Organism", "Treatment", "Temperature_C", "Time_h"], as_index=False)["Zone_mm"]
        .mean()
        .rename(columns={"Temperature_C": "Temp (°C)", "Time_h": "Time (h)", "Zone_mm": "Mean Zone (mm)"})
    )
    add_df_table(
        doc,
        anti_means,
        "Table 5. Mean Zone of Inhibition by Treatment and Organism",
        note="Zone of inhibition values are reported in millimeters. The E. coli workbook contained one blank replicate that was retained as zero to match workbook formulas.",
        digits=3,
        font_size=7.5,
    )
    doc.add_picture(str(FIG_DIR / "antimicrobial_zone_by_treatment.png"), width=Inches(6.25))
    add_caption(doc, "Figure 3. Zone of inhibition by treatment and test organism.")

    anti_anova_out = anti_anova.rename(columns={"p_value": "p-value", "Significant_0.05": "Significant at 0.05"})
    add_df_table(doc, anti_anova_out, "Table 6. One-Way ANOVA for Antimicrobial Activity", font_size=8.5)

    doc.add_heading("Response Surface Methodology and Optimization", level=1)
    doc.add_paragraph(
        "Second-order response surface models were fitted to treatment means to estimate response trends across the experimental region. The composite desirability approach favored higher WAC, WSI, and antimicrobial zone values, while favoring lower water activity and moisture content for product stability."
    )
    doc.add_picture(str(FIG_DIR / "rsm_contour_surfaces.png"), width=Inches(6.25))
    add_caption(doc, "Figure 4. Response surface contour plots for selected physicochemical and antimicrobial responses.")
    doc.add_picture(str(FIG_DIR / "rsm_composite_desirability.png"), width=Inches(5.7))
    add_caption(doc, "Figure 5. Composite desirability surface showing the best screened processing region.")

    opt_out = opt.head(5).rename(
        columns={
            "Temperature_C": "Temp (°C)",
            "Time_h": "Time (h)",
            "MC_percent": "MC (%)",
            "Zone_Mean_All": "Mean Zone",
            "Composite_Desirability": "Composite D",
        }
    )[["Temp (°C)", "Time (h)", "WAC", "Aw", "MC (%)", "WSI", "Mean Zone", "Composite D"]]
    add_df_table(doc, opt_out, "Table 7. Top RSM Composite Desirability Candidates", digits=3)

    doc.add_heading("Interpretation by Objective", level=1)
    doc.add_heading("Objective 1: Develop Garlic Peel Powder Treatments", level=2)
    doc.add_paragraph(
        "The treatment matrix was successfully structured according to drying temperature and drying time. This enabled direct comparison of nine drying conditions and supported the factorial analysis required by the method."
    )
    doc.add_heading("Objective 2: Determine Physicochemical Characteristics", level=2)
    doc.add_paragraph(
        "The physicochemical objective was achieved. Temperature, time, and their interaction showed statistically significant effects for many responses, particularly color values, moisture content, water activity, and WSI. These results indicate that drying conditions materially influenced powder quality and stability-related properties."
    )
    doc.add_heading("Objective 3: Evaluate Antimicrobial Capacity", level=2)
    doc.add_paragraph(
        "The antimicrobial objective was achieved within the limitations of the recorded data. Staphylococcus aureus showed significant differences among treatments, Escherichia coli did not show significant differences, and Salmonella spp. values were constant at 6 mm, preventing meaningful treatment discrimination."
    )
    doc.add_heading("Objective 4: Identify Optimized Conditions through RSM", level=2)
    doc.add_paragraph(
        "The RSM objective was achieved. The best observed condition was Treatment 5 at 60°C for 6 hours. The model-based desirability screen suggested an optimized region near 57.33°C and 5.87 hours, which is close to the observed Treatment 5 condition."
    )

    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph(
        "Overall, the objectives of the study were achieved computationally using the available data. The analysis supports the use of oven-drying conditions near 60°C for 6 hours as the most balanced observed processing condition for Ilocos white garlic peel powder, considering physicochemical quality and antimicrobial performance together."
    )
    add_bullets(
        doc,
        [
            "Treatment 5 is recommended as the best observed treatment from the current dataset.",
            "The RSM optimum should be validated experimentally because it is model-predicted rather than directly observed.",
            "The manuscript should resolve the Treatment 9 time discrepancy before final submission.",
            "Future runs should avoid blank antimicrobial cells and should consider confirmatory assays around the RSM optimum.",
        ],
    )

    doc.add_heading("Appendix: Generated Analysis Files", level=1)
    appendix_df = pd.DataFrame(
        [
            ["rhodea_garlic_peel_analysis.ipynb", "Executed Python/Jupyter notebook with computations and visualizations."],
            ["outputs/two_way_anova_physicochemical.csv", "Two-way ANOVA table for physicochemical responses."],
            ["outputs/one_way_anova_antimicrobial.csv", "One-way ANOVA table for antimicrobial activity."],
            ["outputs/rsm_optimization_top_candidates.csv", "Top RSM desirability candidates."],
            ["figures/*.png", "Visualizations included in this report."],
        ],
        columns=["File", "Description"],
    )
    add_df_table(doc, appendix_df, "Table 8. Reproducible Analysis Artifacts", font_size=8)

    doc.save(REPORT_PATH)
    print(REPORT_PATH)


if __name__ == "__main__":
    main()
