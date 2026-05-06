from __future__ import annotations

from pathlib import Path
import zipfile
from xml.etree import ElementTree as ET

from docx import Document


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "Garlic_Peel_Powder_Full_Report_Rosal.docx"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def main() -> None:
    doc = Document(REPORT)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = len(doc.tables)
    with zipfile.ZipFile(REPORT) as zf:
        media = [name for name in zf.namelist() if name.startswith("word/media/")]
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    headings = []
    for p in root.findall(".//w:p", NS):
        texts = [t.text for t in p.findall(".//w:t", NS) if t.text]
        text = "".join(texts).strip()
        pstyle = p.find(".//w:pStyle", NS)
        if text and pstyle is not None:
            style = pstyle.attrib.get(f"{{{NS['w']}}}val", "")
            if style.startswith("Heading"):
                headings.append(text)

    required = [
        "FULL ANALYSIS REPORT",
        "Engr. Jamie Eduardo Rosal, MSCpE",
        "Executive Summary",
        "Objective Alignment",
        "Physicochemical Results",
        "Antimicrobial Results",
        "Response Surface Methodology and Optimization",
        "Conclusion",
    ]
    missing = [item for item in required if not any(item in p for p in paragraphs)]

    print(f"Report: {REPORT}")
    print(f"Paragraphs: {len(paragraphs)}")
    print(f"Tables: {tables}")
    print(f"Embedded media files: {len(media)}")
    print(f"Headings: {len(headings)}")
    print("Missing required text:", missing)
    if tables < 8:
        raise SystemExit("Expected at least 8 report tables.")
    if len(media) < 5:
        raise SystemExit("Expected at least 5 embedded figures.")
    if missing:
        raise SystemExit("Required report text missing.")
    print("Structural DOCX validation passed.")


if __name__ == "__main__":
    main()
